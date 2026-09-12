"""Asama 10: grup kayitlarini Excel ekiyle e-postalama.

Uc katman ayri ayri sinanir: saf sablon/HTML mantigi (`app/mailsend.py`), depo ve
uclar (sablon CRUD, sira, grup bagi, onizleme, gonderim), son olarak gecici ek
dosyalari ve arayuz kancalari. Gercek Outlook'a hic dokunulmaz: gonderici sahte.
"""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app import db, mailsend
from app import mailsend_repo
from app import repository as repo
from app.mail import fake as mail_fake
from app.mail import send as mail_send
from app.mail.source import unavailable as mail_unavailable
from app.settings_store import MODE_SERVER
from tests.fake_jira import issue

BASE = "https://jira.example.com"

CATALOG = [
    {"id": "summary", "name": "Özet", "schema": {"type": "string"}},
    {"id": "status", "name": "Durum", "schema": {"type": "status"}},
]

COLUMNS = ["issuekey", "summary", "status"]

ORNEK = "ornek@example.com"
IKINCI = "ikinci@example.com"


def setup_group(api_client, conn, name="Filom"):
    repo.store_fields(conn, CATALOG)
    api_client.put("/api/settings", json={"jira.base_url": BASE, "jira.mode": MODE_SERVER})
    group = api_client.post("/api/groups", json={"name": name, "kind": "manual"}).json()["group"]
    api_client.post(f"/api/groups/{group['id']}/items", json={"text": "DEMO-1 DEMO-2"})
    repo.upsert_issues(
        conn,
        [
            issue("DEMO-1", "Alfa kaydı", status={"name": "Açık"}),
            issue("DEMO-2", "Beta kaydı", status={"name": "Kapalı"}),
        ],
    )
    api_client.put(f"/api/groups/{group['id']}", json={"columns": COLUMNS})
    return group


@pytest.fixture
def sender(context):
    """Sahte gonderici: cagri kaydedilir, Outlook'a hic gidilmez."""
    fake = mail_fake.FakeSender()
    context.sender_factory = lambda: fake
    return fake


@pytest.fixture
def temp_mail(tmp_path, monkeypatch):
    """Ek dosyalari testin kendi klasorune duser."""
    folder = tmp_path / "holocron-mail"
    folder.mkdir()
    monkeypatch.setattr(mail_send, "temp_dir", lambda: folder)
    return folder


# --- goc 0010 -----------------------------------------------------------


def test_migration_creates_the_mail_send_tables(conn):
    tables = db.table_names(conn)
    assert "mail_templates" in tables
    assert "mail_sends" in tables
    columns = {row["name"] for row in conn.execute("PRAGMA table_info(groups)")}
    assert "mail_template_id" in columns
    assert db.current_version(conn) >= 10


def test_migration_seeds_the_weekly_template(conn):
    templates = mailsend_repo.list_templates(conn)
    assert [template["name"] for template in templates] == ["Haftalık liste"]
    seed = templates[0]
    assert seed["subject"] == "{grup} — {tarih} ({adet} kayıt)"
    assert "{tablo}" in seed["body"]
    assert seed["attach_excel"] is True
    assert seed["inline_table"] is True


def test_the_seed_is_written_only_once(conn):
    mailsend_repo.delete_template(conn, mailsend_repo.list_templates(conn)[0]["id"])
    db.migrate(conn)
    assert mailsend_repo.list_templates(conn) == []


# --- saf sablon cozumu --------------------------------------------------


def sample_rows():
    return [
        {
            "key": "DEMO-1",
            "url": f"{BASE}/browse/DEMO-1",
            "cells": [
                {"text": "DEMO-1"},
                {"text": "Alfa kaydı"},
                {"text": "Açık"},
            ],
        }
    ]


def sample_columns():
    return [
        {"id": "issuekey", "name": "Anahtar"},
        {"id": "summary", "name": "Özet"},
        {"id": "status", "name": "Durum"},
    ]


def test_all_placeholders_are_resolved():
    group = {"name": "Filom", "jql": "project = DEMO"}
    values = mailsend.values_for(group, sample_rows(), when=date(2026, 9, 12))
    text = mailsend.render_text("{grup} · {tarih} · {adet} · {jql}", values)
    assert text == "Filom · 12.09.2026 · 1 · project = DEMO"


def test_an_unknown_placeholder_stays_empty():
    values = mailsend.values_for({"name": "Filom"}, [])
    assert mailsend.render_text("A{musteri}B", values) == "AB"


def test_the_subject_drops_the_table_token():
    values = mailsend.values_for({"name": "Filom"}, sample_rows(), when=date(2026, 9, 12))
    assert mailsend.render_subject("{grup} {tablo} {tarih}", values) == "Filom 12.09.2026"


def test_plain_text_bodies_become_html_with_line_breaks():
    values = mailsend.values_for({"name": "Filom"}, [])
    html = mailsend.render_body("Merhaba,\n\nAli & Veli <3", values)
    assert "<br>" in html
    assert "&amp;" in html and "&lt;3" in html
    assert html.startswith("<div style=")


def test_a_html_body_is_kept_as_written():
    values = mailsend.values_for({"name": "Filom"}, [])
    html = mailsend.render_body("<p>Merhaba <b>{grup}</b></p>", values)
    assert "<p>Merhaba <b>Filom</b></p>" in html
    assert "&lt;p&gt;" not in html


def test_the_table_carries_headers_borders_and_a_key_link():
    table = mailsend.build_table(sample_rows(), sample_columns())
    assert "border-collapse:collapse" in table
    assert "font-weight:bold" in table
    assert table.count("<th ") == 3
    assert f'<a href="{BASE}/browse/DEMO-1"' in table
    assert "Alfa kaydı" in table
    # Outlook Word ile cizer: sinif secicisi degil, satir ici stil.
    assert "class=" not in table
    assert "<style" not in table


def test_the_table_escapes_cell_text():
    rows = [{"key": "DEMO-1", "url": "", "cells": [{"text": "<script>"}]}]
    table = mailsend.build_table(rows, [{"id": "summary", "name": "Özet"}])
    assert "&lt;script&gt;" in table
    assert "<script>" not in table


def test_render_mail_fills_subject_body_and_addresses():
    template = {
        "name": "Haftalık",
        "to_addresses": f"{ORNEK}, {IKINCI}",
        "cc_addresses": IKINCI,
        "subject": "{grup} — {adet} kayıt",
        "body": "Merhaba,\n\n{tablo}\n\nİyi çalışmalar.",
        "attach_excel": True,
        "inline_table": True,
    }
    result = mailsend.render_mail(
        template, {"name": "Filom"}, sample_rows(), sample_columns(), when=date(2026, 9, 12)
    )
    assert result["subject"] == "Filom — 1 kayıt"
    assert result["to"] == [ORNEK, IKINCI]
    assert result["cc"] == [IKINCI]
    assert "<table" in result["html"]
    assert "İyi çalışmalar." in result["html"]
    assert result["count"] == 1


def test_the_table_can_be_left_out_of_the_body():
    template = {"subject": "x", "body": "Üstü\n{tablo}\nAltı", "inline_table": False}
    result = mailsend.render_mail(template, {"name": "Filom"}, sample_rows(), sample_columns())
    assert "<table" not in result["html"]
    assert "Üstü" in result["html"] and "Altı" in result["html"]


def test_addresses_are_split_and_deduplicated():
    assert mailsend.split_addresses(f"{ORNEK}; {ORNEK.upper()}\n{IKINCI},") == [ORNEK, IKINCI]
    assert mailsend.join_addresses([ORNEK, IKINCI]) == f"{ORNEK}; {IKINCI}"


def test_the_file_name_is_ascii_safe():
    name = mailsend.file_name("Çağrı Filosu", date(2026, 9, 12))
    assert name == "Cagri-Filosu-2026-09-12.xlsx"
    assert name.isascii()


# --- sablon uclari ------------------------------------------------------


def test_templates_can_be_listed_created_updated_and_deleted(api_client):
    listing = api_client.get("/api/mail-templates").json()
    assert listing["mode"] == "display"
    assert any(item["token"] == "{tablo}" for item in listing["placeholders"])
    assert len(listing["templates"]) == 1

    created = api_client.post(
        "/api/mail-templates",
        json={
            "name": "Günlük",
            "to_addresses": ORNEK,
            "cc_addresses": IKINCI,
            "subject": "{grup}",
            "body": "{tablo}",
            "attach_excel": False,
        },
    ).json()["template"]
    assert created["name"] == "Günlük"
    assert created["attach_excel"] is False
    assert created["inline_table"] is True

    updated = api_client.put(
        f"/api/mail-templates/{created['id']}", json={"subject": "Yeni  konu"}
    ).json()["template"]
    assert updated["subject"] == "Yeni konu"

    assert api_client.get(f"/api/mail-templates/{created['id']}").json()["template"]["id"] == (
        created["id"]
    )
    assert api_client.delete(f"/api/mail-templates/{created['id']}").json()["ok"] is True
    assert api_client.get(f"/api/mail-templates/{created['id']}").status_code == 404


def test_a_template_needs_a_name(api_client):
    response = api_client.post("/api/mail-templates", json={"name": "  "})
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "invalid_name"


def test_templates_can_be_reordered(api_client):
    second = api_client.post(
        "/api/mail-templates", json={"name": "İkinci", "body": "x"}
    ).json()["template"]
    first = api_client.get("/api/mail-templates").json()["templates"][0]

    ordered = api_client.post(
        "/api/mail-templates/reorder", json={"ids": [second["id"], first["id"]]}
    ).json()["templates"]
    assert [item["id"] for item in ordered] == [second["id"], first["id"]]
    assert api_client.get("/api/mail-templates").json()["templates"][0]["id"] == second["id"]


def test_a_group_can_have_a_default_template(api_client, conn):
    group = setup_group(api_client, conn)
    second = api_client.post(
        "/api/mail-templates", json={"name": "İkinci", "subject": "İK", "body": "x"}
    ).json()["template"]

    bound = api_client.put(
        f"/api/groups/{group['id']}/mail-template", json={"template_id": second["id"]}
    ).json()
    assert bound["template_id"] == second["id"]

    preview = api_client.post(f"/api/groups/{group['id']}/mail-preview", json={}).json()
    assert preview["template_id"] == second["id"]

    api_client.put(f"/api/groups/{group['id']}/mail-template", json={"template_id": None})
    again = api_client.post(f"/api/groups/{group['id']}/mail-preview", json={}).json()
    assert again["template_id"] != second["id"]


def test_deleting_a_template_unbinds_the_group(api_client, conn):
    group = setup_group(api_client, conn)
    template = api_client.get("/api/mail-templates").json()["templates"][0]
    api_client.put(f"/api/groups/{group['id']}/mail-template", json={"template_id": template["id"]})
    api_client.delete(f"/api/mail-templates/{template['id']}")
    assert mailsend_repo.group_template_id(conn, group["id"]) is None


# --- onizleme -----------------------------------------------------------


def test_preview_renders_the_group_without_saving_anything(api_client, conn):
    group = setup_group(api_client, conn)
    preview = api_client.post(f"/api/groups/{group['id']}/mail-preview", json={}).json()
    assert preview["count"] == 2
    assert preview["file_name"].endswith(".xlsx")
    assert "Filom" in preview["subject"]
    assert "DEMO-1" in preview["html"] and "DEMO-2" in preview["html"]
    assert api_client.get(f"/api/groups/{group['id']}/mail-sends").json()["sends"] == []


def test_preview_can_be_limited_to_the_selected_keys(api_client, conn):
    group = setup_group(api_client, conn)
    preview = api_client.post(
        f"/api/groups/{group['id']}/mail-preview", json={"keys": ["DEMO-2"]}
    ).json()
    assert preview["count"] == 1
    assert "DEMO-2" in preview["html"]
    assert "DEMO-1" not in preview["html"]


def test_preview_honours_the_visible_filter(api_client, conn):
    group = setup_group(api_client, conn)
    filtered = api_client.post(
        f"/api/groups/{group['id']}/mail-preview", json={"q": "Beta", "apply_view": True}
    ).json()
    assert filtered["count"] == 1
    ignored = api_client.post(
        f"/api/groups/{group['id']}/mail-preview", json={"q": "Beta", "apply_view": False}
    ).json()
    assert ignored["count"] == 2


def test_preview_uses_the_edited_subject_and_body(api_client, conn):
    group = setup_group(api_client, conn)
    preview = api_client.post(
        f"/api/groups/{group['id']}/mail-preview",
        json={"subject": "Elle {adet}", "body": "Yalnız metin"},
    ).json()
    assert preview["subject"] == "Elle 2"
    assert "Yalnız metin" in preview["html"]
    assert "<table" not in preview["html"]


def test_preview_can_pick_the_columns(api_client, conn):
    group = setup_group(api_client, conn)
    preview = api_client.post(
        f"/api/groups/{group['id']}/mail-preview", json={"columns": ["issuekey"]}
    ).json()
    assert preview["html"].count("<th ") == 1
    assert "Alfa kaydı" not in preview["html"]


def test_preview_for_an_unknown_group_is_404(api_client):
    assert api_client.post("/api/groups/999/mail-preview", json={}).status_code == 404


# --- gonderim -----------------------------------------------------------


def send_payload(**extra):
    payload = {"to": [ORNEK], "cc": [IKINCI], "mode": "display"}
    payload.update(extra)
    return payload


def test_sending_builds_the_excel_opens_outlook_and_records_the_send(
    api_client, conn, sender, temp_mail
):
    group = setup_group(api_client, conn)
    result = api_client.post(
        f"/api/groups/{group['id']}/mail-send", json=send_payload()
    ).json()

    assert result["ok"] is True
    assert result["mode"] == "display"
    assert result["count"] == 2
    assert result["file_name"].endswith(".xlsx")

    assert len(sender.sent) == 1
    call = sender.last
    assert call["to"] == [ORNEK]
    assert call["cc"] == [IKINCI]
    assert call["mode"] == "display"
    assert len(call["attachments"]) == 1

    written = Path(call["attachments"][0])
    assert written.parent == temp_mail
    book = load_workbook(io.BytesIO(written.read_bytes()))
    assert book.active.max_row == 3  # baslik + iki kayit

    sends = api_client.get(f"/api/groups/{group['id']}/mail-sends").json()["sends"]
    assert len(sends) == 1
    assert sends[0]["to_text"] == ORNEK
    assert sends[0]["issue_count"] == 2
    assert sends[0]["mode"] == "display"
    assert sends[0]["to_count"] == 1


def test_send_mode_is_recorded_separately(api_client, conn, sender, temp_mail):
    group = setup_group(api_client, conn)
    result = api_client.post(
        f"/api/groups/{group['id']}/mail-send", json=send_payload(mode="send")
    ).json()
    assert result["mode"] == "send"
    assert sender.last["mode"] == "send"
    assert api_client.get(f"/api/groups/{group['id']}/mail-sends").json()["sends"][0]["mode"] == (
        "send"
    )


def test_the_mode_falls_back_to_the_setting(api_client, conn, sender, temp_mail):
    group = setup_group(api_client, conn)
    api_client.put("/api/settings", json={"mailsend.mode": "send"})
    payload = send_payload()
    payload.pop("mode")
    result = api_client.post(f"/api/groups/{group['id']}/mail-send", json=payload).json()
    assert result["mode"] == "send"


def test_only_the_selected_keys_land_in_the_excel(api_client, conn, sender, temp_mail):
    group = setup_group(api_client, conn)
    api_client.post(
        f"/api/groups/{group['id']}/mail-send", json=send_payload(keys=["DEMO-2"])
    )
    written = Path(sender.last["attachments"][0])
    sheet = load_workbook(io.BytesIO(written.read_bytes())).active
    keys = [sheet.cell(row=line, column=1).value for line in range(2, sheet.max_row + 1)]
    assert keys == ["DEMO-2"]


def test_sending_without_the_excel_attaches_nothing(api_client, conn, sender, temp_mail):
    group = setup_group(api_client, conn)
    result = api_client.post(
        f"/api/groups/{group['id']}/mail-send", json=send_payload(attach_excel=False)
    ).json()
    assert result["file_name"] == ""
    assert sender.last["attachments"] == []
    assert list(temp_mail.iterdir()) == []


def test_sending_without_a_recipient_is_refused(api_client, conn, sender, temp_mail):
    group = setup_group(api_client, conn)
    response = api_client.post(
        f"/api/groups/{group['id']}/mail-send", json=send_payload(to=[])
    )
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "mail_no_recipient"
    assert sender.sent == []


def test_sending_an_empty_selection_is_refused(api_client, conn, sender, temp_mail):
    group = setup_group(api_client, conn)
    response = api_client.post(
        f"/api/groups/{group['id']}/mail-send", json=send_payload(q="yok", apply_view=True)
    )
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "mail_no_rows"


def test_the_edited_html_is_sent_as_written(api_client, conn, sender, temp_mail):
    group = setup_group(api_client, conn)
    api_client.post(
        f"/api/groups/{group['id']}/mail-send",
        json=send_payload(html="<div>Elle yazıldı</div>"),
    )
    assert sender.last["html"] == "<div>Elle yazıldı</div>"


def test_sending_is_unavailable_without_outlook(api_client, conn, context, temp_mail):
    """Windows disinda uc dosya bile uretmeden `feature_unavailable` doner."""
    group = setup_group(api_client, conn)

    def missing():
        raise mail_unavailable()

    context.sender_factory = missing
    response = api_client.post(f"/api/groups/{group['id']}/mail-send", json=send_payload())
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "feature_unavailable"
    assert list(temp_mail.iterdir()) == []
    assert api_client.get(f"/api/groups/{group['id']}/mail-sends").json()["sends"] == []


def test_an_outlook_failure_is_reported_and_nothing_is_recorded(api_client, conn, context, temp_mail):
    from app.mail import MailError

    group = setup_group(api_client, conn)
    broken = mail_fake.FakeSender(
        error=MailError("outlook_unavailable", "Outlook'a bağlanılamadı. Outlook açık mı?")
    )
    context.sender_factory = lambda: broken
    response = api_client.post(f"/api/groups/{group['id']}/mail-send", json=send_payload())
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "outlook_unavailable"
    assert api_client.get(f"/api/groups/{group['id']}/mail-sends").json()["sends"] == []


# --- gecici ek dosyalari ------------------------------------------------


def test_the_temp_folder_keeps_at_most_twenty_files(tmp_path):
    folder = tmp_path / "holocron-mail"
    folder.mkdir()
    for index in range(25):
        mail_send.write_attachment(f"dosya-{index:02d}.xlsx", b"x", folder=folder)
    files = sorted(item.name for item in folder.iterdir())
    assert len(files) == mail_send.KEEP_FILES
    # En yenileri kalir: silinen ilk bes dosyadir.
    assert "dosya-24.xlsx" in files
    assert "dosya-00.xlsx" not in files


def test_attachment_names_are_ascii_safe(tmp_path):
    path = mail_send.write_attachment("Çağrı Filosu.xlsx", b"x", folder=tmp_path)
    assert path.name == "Cagri-Filosu.xlsx"
    assert path.name.isascii()


def test_the_attachment_is_not_deleted_after_sending(api_client, conn, sender, temp_mail):
    group = setup_group(api_client, conn)
    api_client.post(f"/api/groups/{group['id']}/mail-send", json=send_payload())
    written = Path(sender.last["attachments"][0])
    assert written.exists(), "Outlook eklemeyi asenkron yapabilir: dosya silinmemeli"


def test_the_mode_is_cleaned(monkeypatch):
    assert mail_send.clean_mode("send") == "send"
    assert mail_send.clean_mode("SEND") == "send"
    assert mail_send.clean_mode("") == "display"
    assert mail_send.clean_mode("saçma") == "display"


def test_the_outlook_sender_refuses_to_start_off_windows():
    with pytest.raises(Exception) as caught:
        mail_send.OutlookSender()
    assert getattr(caught.value, "code", "") == "feature_unavailable"


# --- secili kayitlarla Excel --------------------------------------------


def test_export_selected_writes_only_the_given_keys(api_client, conn):
    group = setup_group(api_client, conn)
    response = api_client.get(
        f"/api/groups/{group['id']}/export-selected.xlsx?keys=DEMO-2&columns=issuekey,summary"
    )
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet.max_row == 2
    assert sheet.cell(row=2, column=1).value == "DEMO-2"


def test_export_selected_without_keys_writes_the_whole_group(api_client, conn):
    group = setup_group(api_client, conn)
    response = api_client.get(f"/api/groups/{group['id']}/export-selected.xlsx")
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet.max_row == 3


# --- arayuz kancalari ---------------------------------------------------


def test_the_toolbar_carries_the_button_and_the_selection_badge(api_client):
    page = api_client.get("/").text
    for marker in ('id="mail-send"', "E-posta ile gönder", 'id="pick-count"',
                   "/static/js/mailsend.js"):
        assert marker in page, marker


def test_the_grid_gets_a_checkbox_column(api_client):
    script = api_client.get("/static/js/app.js").text
    for marker in (
        "selectedKeys: new Set()",
        "function clearSelection",
        "function toggleAllVisible",
        "function renderPickCount",
        'id: "pick-all"',
        "Görünen satırların tümünü seç",
        "Bu kaydı seç",
        "seçili",
        "export-selected.xlsx",
        "Yalnız seçili",
    ):
        assert marker in script, marker
    css = api_client.get("/static/css/app.css").text
    assert ".grid td.pick" in css
    assert ".pick-count" in css


def test_the_send_modal_hooks_are_in_its_own_script(api_client):
    script = api_client.get("/static/js/mailsend.js").text
    for marker in (
        "function mailSendModal",
        "/api/mail-templates",
        "mail-preview",
        "mail-send",
        "mail-sends",
        'id="mailsend-template"' .replace('id="', "").replace('"', ""),
        "mailsend-subject",
        "mailsend-body",
        "mailsend-columns",
        "mailsend-preview",
        "mailsend-history",
        "Outlook'ta aç",
        "Gönderilenler",
        "Excel ekle",
        "Gövdeye tablo ekle",
        "Görünen süzgeç ve sıralamayı uygula",
        "{tablo}",
    ):
        assert marker in script, marker


def test_the_settings_card_manages_the_templates(api_client):
    page = api_client.get("/settings").text
    for marker in (
        'id="mailsend-card"',
        "E-posta şablonları",
        'id="mailsend-templates"',
        'id="mailsend-name"',
        'id="mailsend-new-body"',
        'id="mailsend-mode"',
        'value="send"',
        "Doğrudan gönder",
        "Outlook'ta aç",
        "ornek@example.com",
    ):
        assert marker in page, marker
    script = api_client.get("/static/js/mailsend-settings.js").text
    for marker in (
        "/api/mail-templates",
        "mail-templates/reorder",
        "mailsend.mode",
        "function renderMailTemplates",
        "function addMailTemplate",
        "function dropMailTemplate",
    ):
        assert marker in script, marker


def test_the_send_mode_setting_round_trips(api_client):
    assert api_client.get("/api/settings").json()["settings"]["mailsend.mode"] == "display"
    saved = api_client.put("/api/settings", json={"mailsend.mode": "send"}).json()["settings"]
    assert saved["mailsend.mode"] == "send"
    assert api_client.get("/api/mail-templates").json()["mode"] == "send"

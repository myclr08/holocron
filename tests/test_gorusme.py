"""Gorusme notlari: algilama, kayit, kuyruk, ozet ve notun kendisi.

Test edilen sorular:

* Mikrofon izin defterindeki hangi satir "Teams su an konusuyor" der?
* Takip duraklatilmisken kayit basliyor mu? (baslamamali)
* Gorusme icinde aygit degisirse ikinci parca aciliyor mu?
* Asgari surenin altindaki gorusme isleniyor mu? (islenmemeli, "atlandi")
* Kuyruk sirali mi, uygulama yarim isi acilista surduruyor mu?
* Model reddedilirse sradaki deneniyor ve "son calisan" yaziliyor mu?
* faster-whisper yoksa ses siliniyor mu? (SILINMEMELI)
* Copilot vekili alt surecin ortamina yaziliyor, Holocron'unkine yazilmiyor mu?

Gercek ses karti, faster-whisper ve Copilot CLI hicbir testte calismaz:
`app/gorusme/sahte.py` bellek ici karsiliklaridir.
"""

from __future__ import annotations

import json
import os
import subprocess
import sys
import threading
import time
import wave
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest

from app import db, repository as repo
from app.gorusme import algilayici as algilayici_modulu
from app.gorusme import birlestir, depo, intake as gorusme_intake, ozet, sahte, yaziyadok
from app.gorusme import kayit as kayit_modulu
from app.gorusme.algilayici import DefterSatiri, KayitDefteriAlgilayici, mikrofon_acik, teams_mi
from app.gorusme.kuyruk import Kuyruk
from app.gorusme.source import (
    DURUM_ATLANDI,
    DURUM_HATA,
    DURUM_HAZIR,
    DURUM_KUYRUKTA,
    DURUM_YAZIYA_DOKULUYOR,
    KANAL_HOP,
    KANAL_MIK,
    GorusmeHatasi,
    OzetCikti,
    Parca,
    Segment,
    calisma_koku,
)
from app.gorusme.takipci import Takipci
from app.teamscalls.fake import ME, PERSON_ONE, call

NOW = datetime(2026, 9, 18, 14, 10, tzinfo=timezone.utc)


class Saat:
    """Testlerin ilerlettigi saat: hicbir test gercekten beklemez."""

    def __init__(self, baslangic: datetime = NOW) -> None:
        self.simdi = baslangic

    def __call__(self) -> datetime:
        return self.simdi

    def ilerle(self, saniye: float) -> None:
        self.simdi = self.simdi + timedelta(seconds=saniye)


def kur_takipci(context, fake_gorusme, saat: Saat, kuyruk=None) -> Takipci:
    return Takipci(
        context,
        algilayici=fake_gorusme["algilayici"],
        kayitci=fake_gorusme["kayitci"],
        kuyruk=kuyruk,
        bildirimci=fake_gorusme["bildirimci"],
        saat=saat,
    )


def kur_kuyruk(context, fake_gorusme, gorusme_suruyor=None) -> Kuyruk:
    return Kuyruk(
        context,
        dokucu_factory=lambda ayarlar: fake_gorusme["dokucu"],
        ozetleyici_factory=lambda ayarlar: fake_gorusme["ozetleyici"],
        bildirimci=fake_gorusme["bildirimci"],
        gorusme_suruyor=gorusme_suruyor,
    )


def takibi_ac(context, min_dakika: str = "4") -> None:
    context.settings.set("calls.takip", "1")
    context.settings.set("calls.min_dakika", min_dakika)


# --- Asama A: algilama ---------------------------------------------------


def test_registry_key_names_are_recognised_for_both_teams_flavours():
    assert teams_mi("MSTeams_8wekyb3d8bbwe!MSTeams")
    assert teams_mi(r"C:#Program Files#WindowsApps#MSTeams#ms-teams.exe")
    assert teams_mi("C:#Users#ornek#AppData#Local#Microsoft#Teams#current#Teams.exe")
    assert not teams_mi("Microsoft.WindowsCamera_8wekyb3d8bbwe")
    assert not teams_mi("")


def test_microphone_is_in_use_only_when_stop_is_zero():
    satirlar = [
        DefterSatiri("MSTeams_8wekyb3d8bbwe!MSTeams", baslangic=133, bitis=0),
        DefterSatiri("Microsoft.WindowsCamera_8wekyb3d8bbwe", baslangic=1, bitis=0),
    ]
    assert mikrofon_acik(satirlar)
    # Teams mikrofonu birakmis: Stop dolu.
    kapali = [DefterSatiri("MSTeams_8wekyb3d8bbwe!MSTeams", baslangic=133, bitis=140)]
    assert not mikrofon_acik(kapali)
    # Baslangic damgasi yoksa "kullanimda" denmez.
    assert not mikrofon_acik([DefterSatiri("MSTeams_8wekyb3d8bbwe", baslangic=0, bitis=0)])


def test_detector_falls_back_to_the_audio_session_and_never_dies_on_it():
    cagrilar: list[str] = []

    def patlayan() -> bool:
        cagrilar.append("oturum")
        raise RuntimeError("comtypes patladi")

    dedektor = KayitDefteriAlgilayici(
        okuyucu=lambda: [], aygit_cozucu=lambda: sahte.AYGITLAR, oturum_sinyali=patlayan
    )
    assert dedektor.durum().aktif is False
    assert cagrilar == ["oturum"]

    calisan = KayitDefteriAlgilayici(
        okuyucu=lambda: [], aygit_cozucu=lambda: sahte.AYGITLAR, oturum_sinyali=lambda: True
    )
    durum = calisan.durum()
    assert durum.aktif is True
    assert durum.kaynak == "ses_oturumu"
    assert durum.aygitlar.mikrofon == sahte.AYGITLAR.mikrofon


@pytest.mark.skipif(os.name == "nt", reason="Windows'ta winreg zaten var")
def test_registry_reader_never_touches_windows_api_on_linux():
    """Linux'ta `winreg` hic ice aktarilmaz; okuma bos liste doner."""
    assert algilayici_modulu._winreg() is None
    assert algilayici_modulu.defteri_oku() == []


# --- Asama A: kayit ve takip ---------------------------------------------


def test_paused_tracking_records_nothing(context, fake_gorusme):
    saat = Saat()
    takipci = kur_takipci(context, fake_gorusme, saat)
    fake_gorusme["algilayici"].basla()
    context.settings.set("calls.takip", "0")

    assert takipci.yokla() is None
    assert fake_gorusme["kayitci"].basla_cagrisi == 0
    assert depo.list_notlar(context.conn) == []


def test_a_call_is_recorded_and_queued_when_long_enough(context, fake_gorusme):
    saat = Saat()
    takipci = kur_takipci(context, fake_gorusme, saat)
    takibi_ac(context)

    fake_gorusme["algilayici"].basla()
    ozet_satiri = takipci.yokla()
    assert ozet_satiri is not None
    not_id = ozet_satiri["id"]
    kayit = depo.require_not(context.conn, not_id)
    assert kayit["durum"] == "kaydediliyor"
    assert Path(kayit["klasor"]).exists()
    # Kayit basladi bildirimi dustu.
    assert fake_gorusme["bildirimci"].mesajlar[0][0] == "Görüşme kaydediliyor"

    saat.ilerle(27 * 60)
    fake_gorusme["algilayici"].bitir()
    assert takipci.yokla() is None

    kayit = depo.require_not(context.conn, not_id)
    assert kayit["durum"] == DURUM_KUYRUKTA
    assert kayit["sure_sn"] == 27 * 60


def test_a_short_call_is_skipped_and_its_folder_is_deleted(context, fake_gorusme):
    saat = Saat()
    takipci = kur_takipci(context, fake_gorusme, saat)
    takibi_ac(context)

    fake_gorusme["algilayici"].basla()
    not_id = takipci.yokla()["id"]
    klasor = Path(depo.require_not(context.conn, not_id)["klasor"])
    assert klasor.exists()

    # Dort dakikanin altinda: islenmez.
    saat.ilerle(90)
    fake_gorusme["algilayici"].bitir()
    takipci.yokla()

    kayit = depo.require_not(context.conn, not_id)
    assert kayit["durum"] == DURUM_ATLANDI
    assert not klasor.exists()


def test_a_device_change_opens_a_second_part(context, fake_gorusme):
    saat = Saat()
    takipci = kur_takipci(context, fake_gorusme, saat)
    takibi_ac(context)

    fake_gorusme["algilayici"].basla()
    not_id = takipci.yokla()["id"]
    klasor = Path(depo.require_not(context.conn, not_id)["klasor"])
    assert sorted(yol.name for yol in klasor.glob("*.wav")) == ["hop-01.wav", "mik-01.wav"]

    # Kulaklik degisti: yoklama farki gorur, yeni parca acar.
    saat.ilerle(120)
    fake_gorusme["algilayici"].basla(sahte.IKINCI_AYGITLAR)
    takipci.yokla()
    assert sorted(yol.name for yol in klasor.glob("*.wav")) == [
        "hop-01.wav",
        "hop-02.wav",
        "mik-01.wav",
        "mik-02.wav",
    ]


def test_pausing_during_a_call_closes_the_recording(context, fake_gorusme):
    saat = Saat()
    takipci = kur_takipci(context, fake_gorusme, saat)
    takibi_ac(context)
    fake_gorusme["algilayici"].basla()
    not_id = takipci.yokla()["id"]

    saat.ilerle(10 * 60)
    context.settings.set("calls.takip", "0")
    assert takipci.yokla() is None
    assert depo.require_not(context.conn, not_id)["durum"] == DURUM_KUYRUKTA
    assert fake_gorusme["kayitci"].bitir_cagrisi >= 1


def test_the_trial_recording_measures_both_channels(tmp_path):
    kayitci = sahte.SahteKayitci(parca_sn=0.3)
    sonuc = kayitci.deneme(0.3, sahte.AYGITLAR, tmp_path / "deneme")
    assert sonuc["kanallar"][KANAL_MIK]["ses_var"] is True
    assert sonuc["kanallar"][KANAL_HOP]["ses_var"] is True

    sessiz = sahte.SahteKayitci(parca_sn=0.3, sessiz=True)
    bos = sessiz.deneme(0.3, sahte.AYGITLAR, tmp_path / "sessiz")
    assert bos["kanallar"][KANAL_MIK]["ses_var"] is False


# --- Asama A2: dayanikli kayit -------------------------------------------
#
# Saha hatasi (19 Eylul 2026, Windows VDI): dongu aygiti 10 saniye boyunca
# tek cerceve vermedi, WAV 0 bayt kaldi, deneme kaydi ucu `EOFError` ile
# HTTP 500 dondu. Asagidaki testler o zincirin her halkasini tutuyor.


def kayitci_kur(kart: sahte.SahteSesKarti) -> kayit_modulu.WasapiKayitci:
    return kayit_modulu.WasapiKayitci(modul=sahte.SahtePyAudio(kart))


def test_empty_missing_and_broken_wav_files_never_raise(tmp_path):
    bos = tmp_path / "bos.wav"
    bos.write_bytes(b"")
    bozuk = tmp_path / "bozuk.wav"
    bozuk.write_bytes(b"RIFF????BOZUK")
    yok = tmp_path / "yok.wav"
    for yol in (bos, bozuk, yok):
        assert birlestir.oku(yol) == (b"", 16000)
        assert birlestir.dosya_rms(yol) == 0.0
        assert birlestir.sure_sn(yol) == 0.0
        assert birlestir.cerceve_sayisi(yol) == 0
        assert birlestir.bos_mu(yol) is True


def test_the_trial_summary_survives_a_zero_byte_part(tmp_path):
    yol = tmp_path / "mik-00.wav"
    yol.write_bytes(b"")
    ozet_sonuc = kayit_modulu.sonuc_ozeti([Parca(kanal=KANAL_MIK, sira=0, yol=yol)], tmp_path)
    assert ozet_sonuc["kanallar"][KANAL_MIK]["cerceve"] == 0
    assert ozet_sonuc["kanallar"][KANAL_MIK]["ses_var"] is False
    assert ozet_sonuc["kanallar"][KANAL_MIK]["sure_sn"] == 0.0
    assert ozet_sonuc["oneriler"]


def test_the_merge_skips_parts_that_have_no_frames(tmp_path):
    birlestir.yaz(tmp_path / "mik-01.wav", sahte.sentetik_cerceve(0.25))
    birlestir.yaz(tmp_path / "mik-02.wav", b"")  # acildi, veri gelmedi
    (tmp_path / "mik-03.wav").write_bytes(b"")  # basligi bile yazilmamis
    hedef = birlestir.birlestir(sorted(tmp_path.glob("mik-*.wav")), tmp_path / "mik.wav")
    assert hedef is not None
    assert birlestir.sure_sn(hedef) == pytest.approx(0.25, abs=0.02)

    yalnizca_bos = tmp_path / "hop"
    yalnizca_bos.mkdir()
    birlestir.yaz(yalnizca_bos / "hop-01.wav", b"")
    parcalar = list(yalnizca_bos.glob("hop-*.wav"))
    assert birlestir.birlestir(parcalar, yalnizca_bos / "hop.wav") is None


def test_the_recorder_reads_in_callback_mode_and_reports_the_device(tmp_path):
    kart = sahte.SahteSesKarti(bloklar=6)
    sonuc = kayitci_kur(kart).deneme(0.2, sahte.AYGIT_INDEKSLERI, tmp_path)

    mik = sonuc["kanallar"][KANAL_MIK]
    assert mik["acildi"] is True and mik["hata"] == ""
    assert mik["cerceve"] > 0 and mik["ses_var"] is True
    assert mik["aygit"]["ad"] == "Örnek Mikrofon"
    assert mik["aygit"]["hostApi"] == 2
    assert mik["aygit"]["defaultSampleRate"] == 48000
    # Bloklayan `read` YOK: her giris akisi geri cagriyla aciliyor.
    assert all(kwargs.get("stream_callback") for kwargs in kart.giris_acilislari)
    # Dongu aygiti kendi kanal sayisi ve hiziyla aciliyor.
    dongu = kart.giris_acilislari[1]
    assert (dongu["rate"], dongu["channels"]) == (48000, 2)
    # Deneme sirasinda hoparlore duyulmayan sinyal calinir (dongu beslensin).
    assert any(kwargs.get("output") for kwargs in kart.acilislar)
    assert kart.calan_blok > 0
    assert sonuc["oneriler"] == []


def test_a_device_that_sends_nothing_leaves_a_readable_empty_file(tmp_path):
    kart = sahte.SahteSesKarti(bloklar=0)
    sonuc = kayitci_kur(kart).deneme(0.15, sahte.AYGIT_INDEKSLERI, tmp_path)

    assert sonuc["hata"] == ""
    for kanal in (KANAL_MIK, KANAL_HOP):
        veri = sonuc["kanallar"][kanal]
        assert veri["acildi"] is True
        assert veri["cerceve"] == 0
        assert veri["ses_var"] is False
        assert "veri gelmedi" in veri["hata"]
        yol = Path(veri["yol"])
        # Sahadaki 0 bayt dosya bir daha olmasin: baslik her zaman yazilir.
        assert yol.exists() and yol.stat().st_size >= 44
        with wave.open(str(yol), "rb") as dosya:
            assert dosya.getnframes() == 0
    assert "ses çalın" in sonuc["kanallar"][KANAL_HOP]["hata"]
    assert sonuc["oneriler"]


def test_a_part_with_no_frames_is_marked_empty(tmp_path):
    """Gercek kayitta da veri gelmeyen parca "bos" isaretlenir."""
    sessiz = kayitci_kur(sahte.SahteSesKarti(bloklar=0))
    sessiz.basla(tmp_path / "sessiz", 1, sahte.AYGIT_INDEKSLERI)
    parcalar = sessiz.bitir()
    assert parcalar and all(parca.bos for parca in parcalar)

    dolu = kayitci_kur(sahte.SahteSesKarti(bloklar=4))
    dolu.basla(tmp_path / "dolu", 1, sahte.AYGIT_INDEKSLERI)
    time.sleep(0.05)
    assert all(not parca.bos for parca in dolu.bitir())


def test_the_stream_retries_with_the_fallback_format(tmp_path):
    kart = sahte.SahteSesKarti(bloklar=4, acilis_hatasi=1)
    sonuc = kayitci_kur(kart).deneme(0.2, sahte.AYGIT_INDEKSLERI, tmp_path)

    ilk, ikinci = kart.giris_acilislari[0], kart.giris_acilislari[1]
    assert (ilk["rate"], ilk["channels"]) == (48000, 1)
    assert (ikinci["rate"], ikinci["channels"]) == (44100, 2)
    assert sonuc["kanallar"][KANAL_MIK]["acildi"] is True
    assert sonuc["kanallar"][KANAL_MIK]["cerceve"] > 0


def test_a_stuck_writer_thread_still_closes_the_file(tmp_path, monkeypatch):
    """Yazici takilirsa bile dosya kapanir: WAV 0 bayt kalmaz."""
    kart = sahte.SahteSesKarti(bloklar=2)
    yol = tmp_path / "mik-01.wav"
    yazici = kayit_modulu._KanalYazici(kart, 0, yol, False, kanal=KANAL_MIK)
    engel = threading.Event()
    # Takilmis yazici: is parcacigi `bitir`in zaman asimi icinde donmuyor.
    monkeypatch.setattr(yazici, "_yaz_dongusu", lambda: engel.wait(5.0))
    assert yazici.basla() is True
    try:
        yazici.bitir(zaman_asimi=0.05)
        assert "kapanmadı" in yazici.hata
        assert yol.stat().st_size >= 44
        with wave.open(str(yol), "rb") as dosya:
            assert dosya.getnchannels() == 1
            assert dosya.getnframes() == 0
    finally:
        engel.set()


def test_a_channel_that_cannot_be_opened_reports_the_driver_error(tmp_path):
    kart = sahte.SahteSesKarti(bloklar=4, acilis_hatasi=10)
    sonuc = kayitci_kur(kart).deneme(0.1, sahte.AYGIT_INDEKSLERI, tmp_path)
    mik = sonuc["kanallar"][KANAL_MIK]
    assert mik["acildi"] is False
    assert "Invalid sample rate" in mik["hata"]


def test_a_recording_without_audio_fails_with_a_clear_message(context, fake_gorusme):
    saat = Saat()
    not_id = hazirla_not(context, fake_gorusme, saat)
    klasor = Path(depo.require_not(context.conn, not_id)["klasor"])
    # Iki kanal da acildi ama veri gelmedi: parcalar 0 cerceve.
    for ad in ("mik-01.wav", "hop-01.wav"):
        birlestir.yaz(klasor / ad, b"")

    kuyruk = kur_kuyruk(context, fake_gorusme)
    assert kuyruk.sirayi_isle() == 1

    kayit = depo.require_not(context.conn, not_id)
    assert kayit["durum"] == DURUM_HATA
    assert "Ses alınamadı" in kayit["hata"]
    assert "mikrofon: veri gelmedi" in kayit["hata"]
    assert "hoparlör: veri gelmedi" in kayit["hata"]
    # Ses klasoru DURUR: sorun cozulunce yeniden denenebilsin.
    assert klasor.exists()


# --- Asama B: WAV birlestirme --------------------------------------------


def test_parts_are_merged_in_time_order(tmp_path):
    for sira, ton in ((1, 220), (2, 440)):
        birlestir.yaz(tmp_path / f"mik-{sira:02d}.wav", sahte.sentetik_cerceve(0.25, ton))
    hedef = birlestir.birlestir(list(tmp_path.glob("mik-*.wav")), tmp_path / "mik.wav")
    assert hedef is not None
    with wave.open(str(hedef), "rb") as dosya:
        assert dosya.getframerate() == 16000
        assert dosya.getnchannels() == 1
    assert birlestir.sure_sn(hedef) == pytest.approx(0.5, abs=0.02)


def test_merging_nothing_returns_none(tmp_path):
    assert birlestir.birlestir([], tmp_path / "bos.wav") is None


def test_channels_are_labelled_you_and_the_other_side():
    metin = yaziyadok.harmanla(
        [Segment(0.0, 2.0, "Merhaba")], [Segment(3.0, 5.0, "Merhaba, buyurun")]
    )
    assert metin.splitlines() == [
        "[00:00:00] Sen: Merhaba",
        "[00:00:03] Karşı taraf: Merhaba, buyurun",
    ]


# --- Asama B: kuyruk ve hat ----------------------------------------------


def hazirla_not(context, fake_gorusme, saat: Saat, dakika: int = 27) -> int:
    """Kuyruga girmis bir not uretir (kayit basla -> bitir)."""
    takipci = kur_takipci(context, fake_gorusme, saat)
    takibi_ac(context)
    fake_gorusme["algilayici"].basla()
    not_id = takipci.yokla()["id"]
    saat.ilerle(dakika * 60)
    fake_gorusme["algilayici"].bitir()
    takipci.yokla()
    return not_id


def test_the_pipeline_produces_a_ready_note_and_cleans_the_audio(context, fake_gorusme):
    saat = Saat()
    not_id = hazirla_not(context, fake_gorusme, saat)
    klasor = Path(depo.require_not(context.conn, not_id)["klasor"])

    kuyruk = kur_kuyruk(context, fake_gorusme)
    assert kuyruk.sirayi_isle() == 1

    kayit = depo.require_not(context.conn, not_id)
    assert kayit["durum"] == DURUM_HAZIR
    assert kayit["baslik"] == sahte.ORNEK_OZET["baslik"]
    assert kayit["model"] == "claude-sonnet-5"
    # Ses ve ara dosyalar silindi.
    assert not klasor.exists()

    gruplar = gorusme_intake.bolum_gruplari(depo.bolumler(context.conn, not_id))
    assert [satir["metin"] for satir in gruplar["ozet"]] == sahte.ORNEK_OZET["ozet"]
    assert gruplar["aksiyon"][0]["kisi"] == "Ben"
    assert gruplar["aksiyon"][0]["son_tarih"] == "2026-09-25"
    assert [satir["metin"] for satir in gruplar["soru"]] == sahte.ORNEK_OZET["sorular"]
    # Hazir bildirimi notun basligini tasir.
    assert fake_gorusme["bildirimci"].mesajlar[-1] == (
        "Görüşme notu hazır",
        sahte.ORNEK_OZET["baslik"],
    )


def test_the_queue_runs_one_job_at_a_time_in_order(context, fake_gorusme):
    saat = Saat()
    birinci = hazirla_not(context, fake_gorusme, saat)
    saat.ilerle(60)
    ikinci = hazirla_not(context, fake_gorusme, saat)

    kuyruk = kur_kuyruk(context, fake_gorusme)
    assert kuyruk.sirayi_isle(sinir=1) == 1
    assert depo.require_not(context.conn, birinci)["durum"] == DURUM_HAZIR
    assert depo.require_not(context.conn, ikinci)["durum"] == DURUM_KUYRUKTA

    assert kuyruk.sirayi_isle() == 1
    assert depo.require_not(context.conn, ikinci)["durum"] == DURUM_HAZIR


def test_processing_waits_while_a_call_is_running(context, fake_gorusme):
    saat = Saat()
    hazirla_not(context, fake_gorusme, saat)
    context.settings.set("calls.isleme_gorusme_disinda", "1")

    kuyruk = kur_kuyruk(context, fake_gorusme, gorusme_suruyor=lambda: True)
    assert kuyruk.sirayi_isle() == 0

    # Gorusme bitti: is hemen alinir.
    bosta = kur_kuyruk(context, fake_gorusme, gorusme_suruyor=lambda: False)
    assert bosta.sirayi_isle() == 1


def test_half_finished_work_returns_to_the_queue_on_startup(context, fake_gorusme):
    saat = Saat()
    not_id = hazirla_not(context, fake_gorusme, saat)
    depo.durum_yaz(context.conn, not_id, DURUM_YAZIYA_DOKULUYOR)

    assert depo.yarim_isleri_kuyruga_al(context.conn) == 1
    assert depo.require_not(context.conn, not_id)["durum"] == DURUM_KUYRUKTA


def test_a_rejected_model_falls_through_to_the_next_one(context, fake_gorusme):
    saat = Saat()
    not_id = hazirla_not(context, fake_gorusme, saat)
    fake_gorusme["ozetleyici"] = sahte.SahteOzetleyici(reddedilen=["claude-sonnet-5"])

    kuyruk = kur_kuyruk(context, fake_gorusme)
    kuyruk.sirayi_isle()

    kayit = depo.require_not(context.conn, not_id)
    assert kayit["durum"] == DURUM_HAZIR
    assert kayit["model"] == "gpt-5-mini"
    # Son calisan ayara yazildi: sonraki gorusme oradan baslar.
    assert context.settings.get("calls.ozet_model_son") == "gpt-5-mini"
    assert gorusme_intake.load_config(context.settings).model_sirasi()[0] == "gpt-5-mini"


def test_a_failed_stage_keeps_the_audio_and_can_be_retried(context, fake_gorusme):
    saat = Saat()
    not_id = hazirla_not(context, fake_gorusme, saat)
    klasor = Path(depo.require_not(context.conn, not_id)["klasor"])

    class Yok:
        def cevir(self, yol, dil="tr"):
            raise yaziyadok.eksik_paket()

    kuyruk = Kuyruk(
        context,
        dokucu_factory=lambda ayarlar: Yok(),
        ozetleyici_factory=lambda ayarlar: fake_gorusme["ozetleyici"],
        bildirimci=fake_gorusme["bildirimci"],
    )
    kuyruk.sirayi_isle()

    kayit = depo.require_not(context.conn, not_id)
    assert kayit["durum"] == DURUM_HATA
    assert "faster-whisper" in kayit["hata"]
    # Ses SILINMEZ: yeniden denenebilsin.
    assert klasor.exists()

    calisan = kur_kuyruk(context, fake_gorusme)
    calisan.yeniden_dene(not_id)
    assert depo.require_not(context.conn, not_id)["durum"] == DURUM_KUYRUKTA
    calisan.sirayi_isle()
    assert depo.require_not(context.conn, not_id)["durum"] == DURUM_HAZIR


def test_the_transcript_is_only_stored_when_asked(context, fake_gorusme):
    saat = Saat()
    not_id = hazirla_not(context, fake_gorusme, saat)
    kur_kuyruk(context, fake_gorusme).sirayi_isle()
    assert depo.transkript(context.conn, not_id) == ""

    context.settings.set("calls.transkripti_sakla", "1")
    saat.ilerle(3600)
    ikinci = hazirla_not(context, fake_gorusme, saat)
    kur_kuyruk(context, fake_gorusme).sirayi_isle()
    metin = depo.transkript(context.conn, ikinci)
    assert "Sen:" in metin and "Karşı taraf:" in metin


def test_resummarising_needs_a_stored_transcript(context, fake_gorusme):
    saat = Saat()
    not_id = hazirla_not(context, fake_gorusme, saat)
    kuyruk = kur_kuyruk(context, fake_gorusme)
    kuyruk.sirayi_isle()

    with pytest.raises(GorusmeHatasi):
        kuyruk.yeniden_ozetle(not_id)

    depo.transkript_yaz(context.conn, not_id, "[00:00:01] Sen: Deneme")
    fake_gorusme["ozetleyici"] = sahte.SahteOzetleyici(
        cikti={**sahte.ORNEK_OZET, "baslik": "İkinci deneme"}
    )
    kur_kuyruk(context, fake_gorusme).yeniden_ozetle(not_id)
    assert depo.require_not(context.conn, not_id)["baslik"] == "İkinci deneme"


# --- ozet ciktisi ---------------------------------------------------------


def test_json_is_pulled_out_of_free_text():
    govde = '{"baslik": "Deneme", "ozet": ["bir"]}'
    for ham in (
        govde,
        f"İşte not:\n```json\n{govde}\n```\nHazır.",
        f"Buyurun {govde} bitti",
    ):
        assert ozet.ayristir(ham) == {"baslik": "Deneme", "ozet": ["bir"]}
    assert ozet.ayristir("hiç JSON yok") is None
    # Metin icindeki suslu parantez dengeyi bozmamali.
    icice = '{"baslik": "a { b }", "ozet": []}'
    assert ozet.ayristir(f"önsöz {icice}")["baslik"] == "a { b }"


def test_sections_survive_a_sloppy_model_answer():
    veri = {
        "baslik": "  Başlık  ",
        "ozet": "- birinci\n- ikinci",
        "kararlar": ["karar"],
        "aksiyonlar": ["düz metin aksiyon", {"metin": "iş", "kisi": "Ali", "son_tarih": ""}],
        "sorular": [],
    }
    satirlar = ozet.bolumler(veri)
    turler = {satir["tur"] for satir in satirlar}
    assert turler == {"ozet", "karar", "aksiyon"}
    assert ozet.maddeler(veri, "ozet") == ["birinci", "ikinci"]
    assert ozet.aksiyonlar(veri)[1] == {"metin": "iş", "kisi": "Ali", "son_tarih": ""}
    assert ozet.baslik(veri) == "Başlık"


def test_model_order_puts_the_last_working_model_first():
    assert ozet.modelleri_coz('["a", "b", "c"]', son_calisan="c") == ["c", "a", "b"]
    assert ozet.modelleri_coz("", son_calisan="") == list(ozet.VARSAYILAN_MODELLER)
    assert ozet.modelleri_coz("bozuk json [", son_calisan="") == ["bozuk json ["]


def test_the_default_model_is_the_one_working_on_the_users_account():
    # Sahadan 19 Eylul 2026: gpt-5, claude-sonnet-4.5, gpt-4.1 reddedildi;
    # varsayilan sirali liste artik hesapta calisan modelle basliyor.
    assert ozet.VARSAYILAN_MODELLER[0] == "claude-sonnet-5"
    assert "gpt-5-mini" in ozet.VARSAYILAN_MODELLER
    assert "claude-haiku-4.5" in ozet.VARSAYILAN_MODELLER


def test_the_copilot_cli_model_flag_rejection_is_recognized():
    """Sahadan gelen tam metin: `Model "gpt-4.1" from --model flag is not available`."""
    cikti = OzetCikti(kod=1, metin="", hata='Model "gpt-4.1" from --model flag is not available')
    assert ozet.reddedildi_mi(cikti)


def test_rejected_models_are_summarized_on_one_line():
    ozetleyici = sahte.SahteOzetleyici(reddedilen=["gpt-5", "claude-sonnet-4.5", "gpt-4.1"])
    sonuc = ozet.calistir(
        ozetleyici,
        Path("transkript.txt"),
        ["gpt-5", "claude-sonnet-4.5", "gpt-4.1"],
        "istem",
    )
    assert not sonuc.basarili
    assert sonuc.hata == (
        "Copilot modelleri reddetti: gpt-5, claude-sonnet-4.5, gpt-4.1 — "
        "Ayarlar'dan hesabında olan bir model seçin (ör. claude-sonnet-5)."
    )


def test_the_prompt_keeps_the_json_schema_intact(tmp_path):
    istem = ozet.istem_kur(
        ozet.varsayilan_sablon(), tmp_path / "transkript.txt", ["Örnek Kişi"], "2026-09-18", "27 dk"
    )
    assert str(tmp_path / "transkript.txt") in istem
    assert "Örnek Kişi" in istem
    # Sema oldugu gibi durmali: `str.format` kullanilmiyor.
    assert '"aksiyonlar"' in istem and '{"metin"' in istem


# --- Copilot vekili -------------------------------------------------------


def test_the_proxy_only_reaches_the_copilot_subprocess(store, monkeypatch):
    """Vekil alt surecin ortamina yazilir; Holocron'un kendi ortami degismez."""
    monkeypatch.delenv("HTTPS_PROXY", raising=False)
    monkeypatch.delenv("NO_PROXY", raising=False)
    store.set("calls.copilot_proxy", "http://vekil.ornek.local:8080")
    store.set("jira.base_url", "https://jira.ornek.local")
    ayarlar = gorusme_intake.load_config(store)

    ozetleyici = ozet.default_ozetleyici(
        proxy=ayarlar.copilot_proxy, jira_base_url=ayarlar.jira_base_url
    )
    ortam = ozetleyici.ortam()
    assert ortam["HTTPS_PROXY"] == "http://vekil.ornek.local:8080"
    assert ortam["https_proxy"] == "http://vekil.ornek.local:8080"
    # Jira dogrudan goruluyor: alt surecte bile vekile gitmesin.
    assert "jira.ornek.local" in ortam["NO_PROXY"]
    # Holocron'un kendi ortami ASLA degismez.
    assert "HTTPS_PROXY" not in os.environ


def test_the_jira_client_stays_direct_while_the_copilot_proxy_is_set(store):
    from app.jira_client import create_client
    from app.settings_store import PROXY_DIRECT

    store.set("calls.copilot_proxy", "http://vekil.ornek.local:8080")
    store.set("net.proxy_mode", PROXY_DIRECT)
    store.set("jira.base_url", "https://jira.ornek.local")
    client = create_client(store.jira_config())
    assert client.session.proxies == {}
    assert client.session.trust_env is False


def test_an_empty_proxy_leaves_the_environment_alone():
    ortam = ozet.alt_surec_ortami("", "", taban={"PATH": "/usr/bin"})
    assert ortam == {"PATH": "/usr/bin"}


def test_the_connection_test_reports_the_working_model_and_hides_secrets(tmp_path):
    ozetleyici = sahte.SahteOzetleyici(reddedilen=["gpt-5"], kabuk="hazır")
    sonuc = ozet.sina(ozetleyici, ["gpt-5", "claude-sonnet-4.5"], tmp_path)
    assert sonuc["calisiyor"] is True
    assert sonuc["model"] == "claude-sonnet-4.5"
    assert "çalışıyor · model claude-sonnet-4.5" in sonuc["mesaj"]
    # Sinama dosyasi geride kalmaz.
    assert not (tmp_path / ozet.SINAMA_DOSYASI).exists()

    hepsi_red = sahte.SahteOzetleyici(reddedilen=["gpt-5"])
    basarisiz = ozet.sina(hepsi_red, ["gpt-5"], tmp_path)
    assert basarisiz["calisiyor"] is False
    assert basarisiz["mesaj"]


def test_error_text_masks_anything_that_looks_like_a_secret():
    assert "[gizlendi]" in ozet.temizle("hata: Authorization: Bearer abc.def.ghi")
    assert "[gizlendi]" in ozet.temizle("api_key=cok-gizli-bir-deger")
    assert ozet.temizle("") == ""


# --- Copilot'u bulmak -----------------------------------------------------
#
# Saha hatasi (19 Eylul 2026, Windows VDI): kullanici terminalde `copilot`
# calistirabiliyor ama Holocron "Copilot CLI bulunamadi" diyor. Sebep:
# `holocron.bat` uygulamayi `start "" pythonw.exe` ile aciyor, o surec
# kullanicinin guncel PATH'ini gormuyor (npm'in global klasoru cogu kez
# yalnizca kullanici PATH'inde ve o PATH oturumdan sonra degismis).


def _sahte_copilot(klasor: Path, ad: str = "copilot") -> Path:
    klasor.mkdir(parents=True, exist_ok=True)
    yol = klasor / ad
    yol.write_text("@echo off\n", encoding="utf-8")
    yol.chmod(0o755)
    return yol


def test_the_configured_path_wins_over_the_search(tmp_path):
    hedef = _sahte_copilot(tmp_path / "elle", "copilot.exe")
    bulunan = ozet.copilot_bul(str(hedef), ortam={"PATH": str(tmp_path / "bos")})
    assert bulunan == hedef
    # Tirnakli yapistirma da kabul edilir ("where copilot" ciktisi kopyalanir).
    assert ozet.copilot_bul(f'"{hedef}"', ortam={"PATH": str(tmp_path / "bos")}) == hedef


def test_copilot_is_found_in_the_npm_folder_when_the_process_path_is_stale(tmp_path):
    """Ayar bos, PATH'te yok, `%APPDATA%\\npm\\copilot.cmd` var: bulunmali."""
    appdata = tmp_path / "AppData" / "Roaming"
    hedef = _sahte_copilot(appdata / "npm", "copilot.cmd")
    # Uzantisiz node betigi de yaninda durur; Windows'ta calistirilamaz,
    # o yuzden `.cmd` once gelmeli.
    _sahte_copilot(appdata / "npm", "copilot")
    ortam = {"PATH": str(tmp_path / "bos"), "APPDATA": str(appdata)}

    yol, denenen = ozet.copilot_coz("", ortam=ortam, platform="win32")

    assert yol == hedef
    assert "PATH" in denenen


def test_the_user_path_is_read_back_from_the_registry(tmp_path, monkeypatch):
    """Surecin PATH'i bayat: kayit defterindeki kullanici PATH'i taze okunur."""
    hedef = _sahte_copilot(tmp_path / "kayit")
    monkeypatch.setattr(
        ozet, "kayit_defteri_path", lambda platform="": [str(tmp_path / "kayit")]
    )

    yol, denenen = ozet.copilot_coz(
        "", ortam={"PATH": str(tmp_path / "bos")}, platform="win32"
    )

    assert yol == hedef
    assert "kayıt defteri PATH" in denenen


def test_the_error_says_where_it_looked_and_what_to_do(tmp_path):
    yol, denenen = ozet.copilot_coz(
        "", ortam={"PATH": str(tmp_path / "bos")}, platform="win32"
    )
    assert yol is None
    mesaj = ozet.bulunamadi_mesaji(denenen)
    assert "Copilot CLI bulunamadı" in mesaj
    assert "Denenen: PATH" in mesaj
    assert "Copilot yolu" in mesaj and "where copilot" in mesaj


def test_a_cmd_file_is_run_through_cmd_exe_and_the_prompt_stays_off_the_command_line(
    tmp_path,
):
    """`.cmd` dosyasini CreateProcess dogrudan acamaz; cmd.exe araya girer.

    Araya cmd.exe girince komut satiri YENIDEN ayristirilir: istemdeki tirnak,
    `%` ya da `&` komutu parcalar. Bu yuzden istem argüman olarak degil dosya
    olarak gecer.
    """
    yol = _sahte_copilot(tmp_path / "npm", "copilot.cmd")
    istem = 'Oku: {"aksiyonlar": []} & %PATH% | "tırnak"'

    argumanlar, dosya = ozet.komut_kur(yol, "gpt-5", istem, tmp_path, platform="win32")

    assert argumanlar[:5] == ["cmd.exe", "/c", str(yol), "--model", "gpt-5"]
    assert argumanlar[-1] == "--allow-tool=read"
    assert istem not in " ".join(argumanlar)
    assert dosya is not None and dosya.read_text(encoding="utf-8") == istem
    kisa = argumanlar[argumanlar.index("-p") + 1]
    assert ozet.ISTEM_DOSYASI in kisa
    # Kisa yonlendirmede cmd.exe'nin ozel gordugu tek karakter bile yok.
    assert not set('"%&|<>^') & set(kisa)


def test_an_exe_is_run_directly_with_the_prompt_as_an_argument(tmp_path):
    yol = _sahte_copilot(tmp_path / "npm", "copilot.exe")
    argumanlar, dosya = ozet.komut_kur(yol, "gpt-5", "kısa istem", tmp_path, platform="win32")
    assert argumanlar == [
        str(yol),
        "--model",
        "gpt-5",
        "-p",
        "kısa istem",
        "--allow-tool=read",
    ]
    assert dosya is None
    # Windows disinda `.cmd` diye bir sey yok: kabuk dalina hic girilmez.
    assert ozet.kabukla_mi(Path("/usr/bin/copilot.cmd"), platform="linux") is False


def test_the_summariser_resolves_the_path_and_remembers_it(tmp_path, monkeypatch):
    hedef = _sahte_copilot(tmp_path / "npm", "copilot.exe")
    ozetleyici = ozet.CopilotOzetleyici(yol=str(hedef))
    cagrilar: list[list[str]] = []

    class SahteSonuc:
        returncode = 0
        stdout = '{"baslik": "Deneme", "ozet": ["bir"]}'
        stderr = ""

    def calistir(self, argumanlar, klasor):
        cagrilar.append(list(argumanlar))
        return SahteSonuc()

    monkeypatch.setattr(ozet.CopilotOzetleyici, "calistir", calistir)
    transkript = tmp_path / "is" / "transkript.txt"
    transkript.parent.mkdir(parents=True)
    transkript.write_text("metin", encoding="utf-8")

    cikti = ozetleyici.ozetle(transkript, "gpt-5", "istem")

    assert cikti.kod == 0
    assert ozetleyici.son_yol == str(hedef)
    assert cagrilar[0][0] == str(hedef)
    # Istem dosyasi (varsa) geride kalmaz.
    assert not (transkript.parent / ozet.ISTEM_DOSYASI).exists()


def test_the_copilot_subprocess_never_pops_a_console_window_on_windows(tmp_path, monkeypatch):
    """Saha hatasi 19 Eylul 2026: konsolsuz surecte Copilot cagrisi bos bir
    konsol penceresi aciyordu. `sessiz_calistir_ayarlari` gercek `subprocess.run`
    cagrisina gecmeli; Windows dalini Linux'ta da kanitlamak icin STARTUPINFO
    sahte olarak eklenir (gercek Windows'ta zaten var)."""

    class SahteStartupInfo:
        def __init__(self) -> None:
            self.dwFlags = 0
            self.wShowWindow = None

    class SahteSonuc:
        returncode = 0
        stdout = ""
        stderr = ""

    monkeypatch.setattr(sys, "platform", "win32")
    monkeypatch.setattr(subprocess, "STARTUPINFO", SahteStartupInfo, raising=False)
    monkeypatch.setattr(subprocess, "CREATE_NO_WINDOW", 0x08000000, raising=False)
    monkeypatch.setattr(subprocess, "STARTF_USESHOWWINDOW", 1, raising=False)
    monkeypatch.setattr(subprocess, "SW_HIDE", 0, raising=False)

    yakalanan: dict = {}

    def sahte_run(argumanlar, **kwargs):
        yakalanan.update(kwargs)
        return SahteSonuc()

    monkeypatch.setattr(ozet.subprocess, "run", sahte_run)

    ozetleyici = ozet.CopilotOzetleyici()
    sonuc = ozetleyici.calistir(["copilot", "--model", "gpt-5"], tmp_path)

    assert sonuc.returncode == 0
    assert yakalanan["stdin"] == subprocess.DEVNULL
    assert yakalanan["creationflags"] == 0x08000000
    assert isinstance(yakalanan["startupinfo"], SahteStartupInfo)
    assert yakalanan["startupinfo"].wShowWindow == 0


def test_a_missing_copilot_tells_the_user_where_to_look(tmp_path, monkeypatch):
    monkeypatch.setenv("PATH", str(tmp_path / "bos"))
    cikti = ozet.CopilotOzetleyici().ozetle(tmp_path / "transkript.txt", "gpt-5", "istem")
    assert cikti.kod == 127
    assert "Copilot CLI bulunamadı" in cikti.hata
    assert "Denenen" in cikti.hata


def test_the_summary_setting_carries_the_path_into_the_subprocess(store):
    store.set("calls.copilot_yolu", r"C:\Users\ornek\AppData\Roaming\npm\copilot.cmd")
    ayarlar = gorusme_intake.load_config(store)
    assert ayarlar.copilot_yolu.endswith("copilot.cmd")
    ozetleyici = ozet.default_ozetleyici(yol=ayarlar.copilot_yolu)
    assert ozetleyici.yol == ayarlar.copilot_yolu



# --- yaziya dokme paketinin kurulumu --------------------------------------
#
# Saha hatasi (19 Eylul 2026): lite kurulumda paket hic gelmemisti. Artik
# `requirements.txt` icinde, ayrica Ayarlar'dan da kurulabiliyor. Hicbir test
# gercekten pip calistirmaz: kosucu disaridan verilir.


class PipCiktisi:
    def __init__(self, kod: int, metin: str) -> None:
        self.returncode = kod
        self.stdout = metin
        self.stderr = ""


class SahtePip:
    """pip yerine gecen kosucu: cagrilari ve ortamlari toplar."""

    def __init__(self, kodlar, cikti: str = "tamam") -> None:
        self.kodlar = list(kodlar)
        self.cikti = cikti
        self.cagrilar: list[list[str]] = []
        self.ortamlar: list[dict[str, str]] = []

    def __call__(self, komut, **kwargs):
        self.cagrilar.append(list(komut))
        self.ortamlar.append(dict(kwargs.get("env") or {}))
        kod = self.kodlar.pop(0) if self.kodlar else 0
        return PipCiktisi(kod, self.cikti)


def tekerlek_koy(kok: Path, klasor: str, ad: str) -> None:
    (kok / klasor).mkdir(parents=True, exist_ok=True)
    (kok / klasor / ad).write_bytes(b"")


def test_installing_the_transcription_package_prefers_the_local_wheels(tmp_path, monkeypatch):
    monkeypatch.delenv("HTTPS_PROXY", raising=False)
    tekerlek_koy(tmp_path, "wheels", "faster_whisper-1.2.1-py3-none-any.whl")
    tekerlek_koy(tmp_path, "whisper-wheels", "ctranslate2-4.8.2-cp313-cp313-win_amd64.whl")
    pip = SahtePip([0])

    sonuc = yaziyadok.kur(
        proxy="http://vekil.ornek.local:8080", kok=tmp_path, calistirici=pip
    )

    assert sonuc["kuruldu"] is True
    assert sonuc["kaynak"] == "yerel"
    komut = pip.cagrilar[0]
    assert "--no-index" in komut
    assert str(tmp_path / "wheels") in komut
    assert str(tmp_path / "whisper-wheels") in komut
    assert komut[-1] == yaziyadok.PAKET
    # Vekil YALNIZCA alt surecin ortamina yazilir (ozet.alt_surec_ortami kalibi).
    assert pip.ortamlar[0]["HTTPS_PROXY"] == "http://vekil.ornek.local:8080"
    assert "HTTPS_PROXY" not in os.environ
    # Yerel kurulum tuttu: ag hic denenmedi.
    assert len(pip.cagrilar) == 1


def test_installing_falls_back_to_the_network_when_the_wheels_do_not_fit(tmp_path):
    tekerlek_koy(tmp_path, "wheels", "fastapi-0.115.6-py3-none-any.whl")
    pip = SahtePip([1, 0])

    sonuc = yaziyadok.kur(kok=tmp_path, calistirici=pip)

    assert sonuc["kuruldu"] is True
    assert sonuc["kaynak"] == "ağ"
    assert "--no-index" not in pip.cagrilar[1]


def test_installing_without_local_wheels_goes_straight_to_the_network(tmp_path):
    pip = SahtePip([0])
    sonuc = yaziyadok.kur(kok=tmp_path, calistirici=pip)
    assert len(pip.cagrilar) == 1
    assert "--find-links" not in pip.cagrilar[0]
    assert sonuc["kaynak"] == "ağ"


def test_the_pip_install_never_pops_a_console_window_on_windows(tmp_path, monkeypatch):
    """Ayni saha hatasi pip kurulumu icin de gecerli: konsolsuz surecte pip
    cagrisi da bos bir konsol acmamali."""

    class SahteStartupInfo:
        def __init__(self) -> None:
            self.dwFlags = 0
            self.wShowWindow = None

    monkeypatch.setattr(sys, "platform", "win32")
    monkeypatch.setattr(subprocess, "STARTUPINFO", SahteStartupInfo, raising=False)
    monkeypatch.setattr(subprocess, "CREATE_NO_WINDOW", 0x08000000, raising=False)
    monkeypatch.setattr(subprocess, "STARTF_USESHOWWINDOW", 1, raising=False)
    monkeypatch.setattr(subprocess, "SW_HIDE", 0, raising=False)

    yakalanan: dict = {}

    def kosucu(komut, **kwargs):
        yakalanan.update(kwargs)
        return PipCiktisi(0, "tamam")

    sonuc = yaziyadok.kur(kok=tmp_path, calistirici=kosucu)

    assert sonuc["kuruldu"] is True
    assert yakalanan["stdin"] == subprocess.DEVNULL
    assert yakalanan["creationflags"] == 0x08000000
    assert isinstance(yakalanan["startupinfo"], SahteStartupInfo)
    assert yakalanan["startupinfo"].wShowWindow == 0


def test_a_failed_installation_never_leaks_a_secret_to_the_screen(tmp_path):
    pip = SahtePip([1], cikti="pip hatasi: token=cok-gizli-bir-deger")
    sonuc = yaziyadok.kur(kok=tmp_path, calistirici=pip)
    assert sonuc["kuruldu"] is False
    assert "[gizlendi]" in sonuc["cikti"]
    assert "cok-gizli-bir-deger" not in sonuc["mesaj"]
    assert "cok-gizli-bir-deger" not in sonuc["cikti"]


def test_a_missing_pip_is_explained_instead_of_dumped(tmp_path):
    """Tasinabilir tam pakette pip yok: kullaniciya ne yapacagi soylenir."""
    pip = SahtePip([1], cikti="C:\\python-embed\\python.exe: No module named pip")
    sonuc = yaziyadok.kur(kok=tmp_path, calistirici=pip)
    assert sonuc["kuruldu"] is False
    assert "pip" in sonuc["mesaj"]
    assert "paketle birlikte gelir" in sonuc["mesaj"]


# --- paket gelince bekleyen satirlar --------------------------------------


def test_notes_that_failed_without_the_package_return_to_the_queue(context, fake_gorusme):
    not_id = hazirla_not(context, fake_gorusme, Saat())
    depo.hataya_dus(context.conn, not_id, str(yaziyadok.eksik_paket()))
    baska = hazirla_not(context, fake_gorusme, Saat(NOW + timedelta(hours=2)))
    depo.hataya_dus(context.conn, baska, "Özet alınamadı: model yok")

    assert depo.eksik_paket_hatalarini_kuyruga_al(context.conn) == 1

    kayit = depo.require_not(context.conn, not_id)
    assert kayit["durum"] == DURUM_KUYRUKTA
    assert kayit["hata"] == ""
    # Baska sebeple dusen satira dokunulmaz.
    assert depo.require_not(context.conn, baska)["durum"] == DURUM_HATA


def test_the_queue_finishes_those_notes_once_the_package_is_there(context, fake_gorusme):
    not_id = hazirla_not(context, fake_gorusme, Saat())
    klasor = Path(depo.require_not(context.conn, not_id)["klasor"])

    class Yok:
        def cevir(self, yol, dil="tr"):
            raise yaziyadok.eksik_paket()

    eksik = Kuyruk(
        context,
        dokucu_factory=lambda ayarlar: Yok(),
        ozetleyici_factory=lambda ayarlar: fake_gorusme["ozetleyici"],
    )
    eksik.sirayi_isle()
    assert depo.require_not(context.conn, not_id)["durum"] == DURUM_HATA
    # Ses SILINMEZ: is kaldigi yerden surecek.
    assert klasor.exists()

    kuyruk = kur_kuyruk(context, fake_gorusme)
    assert kuyruk.eksik_paket_islerini_kuyruga_al() == 1
    kuyruk.sirayi_isle()
    assert depo.require_not(context.conn, not_id)["durum"] == DURUM_HAZIR


def test_startup_resumes_package_errors_only_when_the_package_is_installed(
    context, fake_gorusme, monkeypatch
):
    from app.gorusme import servis as gorusme_servis

    not_id = hazirla_not(context, fake_gorusme, Saat())
    depo.hataya_dus(context.conn, not_id, str(yaziyadok.eksik_paket()))

    monkeypatch.setattr(gorusme_servis.yaziyadok, "kurulu_mu", lambda: False)
    gorusme_servis.kur(context)
    assert depo.require_not(context.conn, not_id)["durum"] == DURUM_HATA

    monkeypatch.setattr(gorusme_servis.yaziyadok, "kurulu_mu", lambda: True)
    gorusme_servis.kur(context)
    assert depo.require_not(context.conn, not_id)["durum"] == DURUM_KUYRUKTA


# --- yerel model klasoru --------------------------------------------------
#
# Saha hatasi (19 Eylul 2026, VDI + kurumsal vekil): kullanici release'teki
# model zip'ini acti, klasor yolunu ayara yazdi ve yaziya dokme "ConnectTimeout
# ... cannot find the appropriate snapshot folder" diye dustu. Klasor HF
# ONBELLEK KOKU sanilip `download_root`a veriliyordu; faster-whisper ise acilmis
# klasoru DOGRUDAN kabul eder. Asagidaki testler dort dali da sabitler ve
# klasor verildiginde aga CIKILMADIGINI gosterir.


def model_klasoru(kok: Path, *parcalar: str) -> Path:
    """Icinde `model.bin` + `config.json` olan bir model klasoru uretir."""
    klasor = kok.joinpath(*parcalar) if parcalar else kok
    klasor.mkdir(parents=True, exist_ok=True)
    (klasor / "model.bin").write_bytes(b"sahte model")
    (klasor / "config.json").write_text("{}", encoding="utf-8")
    return klasor


class SahteWhisperModel:
    """Kitapligin yerine gecer: hangi argumanla cagrildigini kaydeder."""

    cagrilar: list[tuple[str, dict]] = []
    hata: BaseException | None = None

    def __init__(self, yol, **kwargs):
        SahteWhisperModel.cagrilar.append((str(yol), dict(kwargs)))
        if SahteWhisperModel.hata is not None:
            raise SahteWhisperModel.hata


@pytest.fixture
def sahte_whisper(monkeypatch):
    """`faster_whisper` modulunu sahtesiyle degistirir (paket kurulu olmasa da)."""
    import sys
    import types

    SahteWhisperModel.cagrilar = []
    SahteWhisperModel.hata = None
    modul = types.ModuleType("faster_whisper")
    modul.WhisperModel = SahteWhisperModel
    monkeypatch.setitem(sys.modules, "faster_whisper", modul)
    return SahteWhisperModel


def test_an_unpacked_model_folder_is_passed_to_whisper_directly(tmp_path, sahte_whisper):
    klasor = model_klasoru(tmp_path / "faster-whisper-small")

    yaziyadok.WhisperDokucu("small", str(klasor)).yukle()

    yol, secenekler = sahte_whisper.cagrilar[0]
    assert yol == str(klasor)
    # Klasorun kendisi model: HF onbellegi diye gosterilmez, ag denenmez.
    assert "download_root" not in secenekler
    assert secenekler["device"] == "cpu"
    assert secenekler["compute_type"] == yaziyadok.HESAP_TIPI


@pytest.mark.parametrize("alt", ["small", "faster-whisper-small"])
def test_the_model_can_sit_one_folder_below_the_given_path(tmp_path, sahte_whisper, alt):
    klasor = model_klasoru(tmp_path / "modeller" / alt)

    yaziyadok.WhisperDokucu("small", str(tmp_path / "modeller")).yukle()

    yol, secenekler = sahte_whisper.cagrilar[0]
    assert yol == str(klasor)
    assert "download_root" not in secenekler


def test_a_hugging_face_cache_root_is_used_offline(tmp_path, sahte_whisper):
    model_klasoru(
        tmp_path / "models--Systran--faster-whisper-small" / "snapshots" / "abc123"
    )

    yaziyadok.WhisperDokucu("small", str(tmp_path)).yukle()

    yol, secenekler = sahte_whisper.cagrilar[0]
    assert yol == "small"
    assert secenekler["download_root"] == str(tmp_path)
    # Onbellek duzeninde bile aga CIKILMAZ.
    assert secenekler["local_files_only"] is True


def test_an_empty_model_folder_fails_before_touching_the_network(tmp_path, sahte_whisper):
    (tmp_path / "bos").mkdir()

    with pytest.raises(GorusmeHatasi) as hata:
        yaziyadok.WhisperDokucu("small", str(tmp_path / "bos")).yukle()

    assert hata.value.code == "model_bulunamadi"
    assert "model.bin" in str(hata.value)
    assert "beklenen düzen" in str(hata.value)
    # Kurumsal agda dakikalarca bekletmesin: kitaplik hic cagrilmadi.
    assert sahte_whisper.cagrilar == []


def test_a_missing_model_folder_says_so(tmp_path, sahte_whisper):
    with pytest.raises(GorusmeHatasi) as hata:
        yaziyadok.WhisperDokucu("small", str(tmp_path / "yok")).yukle()

    assert hata.value.code == "model_klasoru_yok"
    assert sahte_whisper.cagrilar == []


def test_an_empty_setting_still_lets_the_library_download(sahte_whisper):
    yaziyadok.WhisperDokucu("small", "").yukle()

    yol, secenekler = sahte_whisper.cagrilar[0]
    assert yol == "small"
    assert "download_root" not in secenekler
    assert "local_files_only" not in secenekler


def test_a_library_failure_becomes_one_short_turkish_line(tmp_path, sahte_whisper):
    model_klasoru(tmp_path)
    sahte_whisper.hata = OSError("ConnectTimeout: " + "x" * 500)

    with pytest.raises(GorusmeHatasi) as hata:
        yaziyadok.WhisperDokucu("small", str(tmp_path)).yukle()

    assert hata.value.code == "model_yuklenemedi"
    mesaj = str(hata.value)
    assert mesaj.startswith("Model yüklenemedi: ")
    assert len(mesaj) <= len("Model yüklenemedi: ") + yaziyadok.HATA_SINIRI


def test_the_model_is_loaded_once_per_transcriber(tmp_path, sahte_whisper):
    model_klasoru(tmp_path)
    dokucu = yaziyadok.WhisperDokucu("small", str(tmp_path))

    dokucu.yukle()
    dokucu.yukle()

    assert len(sahte_whisper.cagrilar) == 1


# --- "Modeli sina" --------------------------------------------------------


def test_testing_the_model_reports_the_resolved_path_and_the_seconds(tmp_path, sahte_whisper):
    klasor = model_klasoru(tmp_path / "faster-whisper-small")

    sonuc = yaziyadok.sina("small", str(klasor))

    assert sonuc["calisiyor"] is True
    assert sonuc["yol"] == str(klasor)
    assert sonuc["mesaj"].startswith(f"hazır: {klasor}, ")
    assert sonuc["mesaj"].endswith(" sn")


def test_testing_the_model_returns_a_clean_error_instead_of_raising(tmp_path, sahte_whisper):
    (tmp_path / "bos").mkdir()

    sonuc = yaziyadok.sina("small", str(tmp_path / "bos"))

    assert sonuc["calisiyor"] is False
    assert sonuc["kod"] == "model_bulunamadi"
    assert "model.bin" in sonuc["mesaj"]
    assert sahte_whisper.cagrilar == []


def test_the_model_test_endpoint_uses_the_unsaved_form_values(api_client, fake_gorusme):
    yanit = api_client.post(
        "/api/gorusme/ayar/model-sina", json={"model": "small", "klasor": ""}
    )

    assert yanit.status_code == 200
    govde = yanit.json()
    assert govde["calisiyor"] is True
    assert govde["yol"] == "sahte-model"
    assert fake_gorusme["dokucu"].yukleme == 1


def test_the_model_test_endpoint_reports_failures_as_text(api_client, fake_gorusme):
    fake_gorusme["dokucu"].yukleme_hatasi = GorusmeHatasi(
        "model_bulunamadi", "Model klasöründe model.bin bulunamadı: D:\\yok"
    )

    govde = api_client.post("/api/gorusme/ayar/model-sina", json={}).json()

    assert govde["calisiyor"] is False
    assert govde["kod"] == "model_bulunamadi"
    assert "model.bin" in govde["mesaj"]


# --- hata yolunda kuyruk durumu -------------------------------------------


def test_a_model_failure_leaves_the_row_in_error_and_clears_the_counter(
    context, fake_gorusme, tmp_path
):
    not_id = hazirla_not(context, fake_gorusme, Saat())
    fake_gorusme["dokucu"].yukleme_hatasi = GorusmeHatasi(
        "model_yuklenemedi", "Model yüklenemedi: ConnectTimeout"
    )

    kur_kuyruk(context, fake_gorusme).sirayi_isle()

    kayit = depo.require_not(context.conn, not_id)
    assert kayit["durum"] == DURUM_HATA
    assert kayit["hata"].startswith("Model yüklenemedi: ")
    # Takip seridindeki "isleniyor n" sayaci ara durumlari sayar: sifirlandi.
    serit = gorusme_intake.serit(context.conn, gorusme_intake.load_config(context.settings))
    assert serit["isleniyor"] == 0
    assert serit["hata"] == 1


def test_a_model_failure_is_not_mistaken_for_a_missing_package(context, fake_gorusme):
    not_id = hazirla_not(context, fake_gorusme, Saat())
    # Hata metninde model klasorunun adi geciyor ("faster-whisper-small"):
    # bu satir paket kurulumu bekleyen satir DEGILDIR, kuyruga donmemeli.
    depo.hataya_dus(
        context.conn,
        not_id,
        "Model yüklenemedi: D:\\modeller\\faster-whisper-small okunamadı",
    )

    assert depo.eksik_paket_hatalarini_kuyruga_al(context.conn) == 0
    assert depo.require_not(context.conn, not_id)["durum"] == DURUM_HATA


def test_a_half_finished_row_returns_to_the_queue_on_startup(context, fake_gorusme):
    from app.gorusme import servis as gorusme_servis

    not_id = hazirla_not(context, fake_gorusme, Saat())
    # Isci model yuklerken oldurulmus: satir "yaziya_dokuluyor"da kalmis.
    depo.durum_yaz(context.conn, not_id, DURUM_YAZIYA_DOKULUYOR)
    serit = gorusme_intake.serit(context.conn, gorusme_intake.load_config(context.settings))
    assert serit["isleniyor"] == 1

    gorusme_servis.kur(context)

    assert depo.require_not(context.conn, not_id)["durum"] == DURUM_KUYRUKTA
    serit = gorusme_intake.serit(context.conn, gorusme_intake.load_config(context.settings))
    assert serit["isleniyor"] == 0
    assert serit["kuyrukta"] == 1


# --- katilimci eslemesi ---------------------------------------------------


def test_a_note_is_matched_to_the_call_that_overlaps_it(context, fake_gorusme):
    saat = Saat()
    not_id = hazirla_not(context, fake_gorusme, saat)

    repo.import_calls(
        context.conn,
        [
            {
                "call_id": "arama-1",
                # Bir dakika kayma: uc dakikalik payin icinde.
                "started_at": (NOW + timedelta(minutes=1)).isoformat().replace("+00:00", "Z"),
                "duration_ms": 27 * 60 * 1000,
                "kind": "one_to_one",
                "counterpart_id": PERSON_ONE,
                "counterpart_name": "Örnek Kişi",
                "participants_json": json.dumps([ME, PERSON_ONE]),
            }
        ],
    )
    aramalar = repo.list_calls(context.conn)
    arama = gorusme_intake.esle_ve_yaz(context.conn, not_id, aramalar, me=ME)

    assert arama is not None
    assert depo.require_not(context.conn, not_id)["call_id"] == "arama-1"
    kisiler = depo.katilimcilar(context.conn, not_id)
    assert [kisi["ad"] for kisi in kisiler] == ["Örnek Kişi"]


def test_a_call_far_away_in_time_or_length_never_matches():
    aramalar = [
        {
            "call_id": "uzak",
            "started_at": (NOW + timedelta(hours=2)).isoformat().replace("+00:00", "Z"),
            "duration_ms": 600000,
        },
        {
            "call_id": "kisa",
            "started_at": NOW.isoformat().replace("+00:00", "Z"),
            "duration_ms": 30000,
        },
    ]
    hedef = NOW.isoformat().replace("+00:00", "Z")
    assert gorusme_intake.eslesen_arama(aramalar, hedef, 27 * 60) is None


def test_a_note_without_participants_says_it_is_waiting(context, fake_gorusme):
    saat = Saat()
    not_id = hazirla_not(context, fake_gorusme, saat)
    kartlar = gorusme_intake.liste(context.conn)
    kart = next(satir for satir in kartlar if satir["id"] == not_id)
    assert kart["katilimcilar"] == []
    assert kart["katilimci_notu"] == "katılımcı bekleniyor"


# --- depo, arama ve gorevler ----------------------------------------------


def test_the_migration_creates_the_meeting_note_tables(conn):
    assert {
        "gorusme_notu",
        "gorusme_bolum",
        "gorusme_katilimci",
        "gorusme_bag",
        "gorusme_transkript",
    } <= db.table_names(conn)
    assert db.SCHEMA_VERSION == 15


def test_notes_can_be_searched_by_their_summary(context, fake_gorusme):
    saat = Saat()
    not_id = hazirla_not(context, fake_gorusme, saat)
    kur_kuyruk(context, fake_gorusme).sirayi_isle()

    bulunan = gorusme_intake.liste(context.conn, q="raporlama")
    assert [kart["id"] for kart in bulunan] == [not_id]
    assert gorusme_intake.liste(context.conn, q="kesinlikle-yok") == []


def test_an_action_becomes_a_task_bound_to_the_note(context, fake_gorusme):
    from app.gorusme.kuyruk import gorev_uret

    saat = Saat()
    not_id = hazirla_not(context, fake_gorusme, saat)
    kur_kuyruk(context, fake_gorusme).sirayi_isle()
    depo.jira_bagla(context.conn, not_id, "PRJ-1432")

    aksiyon = gorusme_intake.bolum_gruplari(depo.bolumler(context.conn, not_id))["aksiyon"][0]
    gorev = gorev_uret(context, not_id, aksiyon["id"])

    assert gorev["title"] == "Geri dönüş planını yaz"
    assert gorev["due_date"] == "2026-09-25"
    assert gorev["issue_key"] == "PRJ-1432"
    assert depo.gorev_kimlikleri(context.conn, not_id) == [gorev["id"]]
    assert depo.gorev_sayilari(context.conn)[not_id] == 1


def test_a_note_is_bound_to_a_single_jira_issue(context, fake_gorusme):
    saat = Saat()
    not_id = hazirla_not(context, fake_gorusme, saat)
    depo.jira_bagla(context.conn, not_id, "PRJ-1")
    depo.jira_bagla(context.conn, not_id, "PRJ-2")
    assert depo.jira_key(context.conn, not_id) == "PRJ-2"
    assert depo.kayit_notlari(context.conn, "PRJ-2") == [not_id]

    depo.jira_bagla(context.conn, not_id, "")
    assert depo.jira_key(context.conn, not_id) == ""


def test_deleting_a_note_takes_its_sections_with_it(context, fake_gorusme):
    saat = Saat()
    not_id = hazirla_not(context, fake_gorusme, saat)
    kur_kuyruk(context, fake_gorusme).sirayi_isle()

    depo.sil(context.conn, not_id)
    assert depo.get_not(context.conn, not_id) is None
    assert depo.bolumler(context.conn, not_id) == []
    assert gorusme_intake.liste(context.conn, q="raporlama") == []


def test_the_strip_counts_what_the_pipeline_is_doing(context, fake_gorusme):
    saat = Saat()
    hazirla_not(context, fake_gorusme, saat)
    ayarlar = gorusme_intake.load_config(context.settings)
    serit = gorusme_intake.serit(context.conn, ayarlar)
    assert serit["kuyrukta"] == 1
    assert serit["takip"] is True
    assert serit["kaydediliyor"] is False
    assert serit["min_dakika"] == 4.0


def test_the_working_folder_never_leaves_the_data_directory(isolated_home):
    assert calisma_koku("") == isolated_home / "gorusme"
    assert calisma_koku(str(isolated_home / "baska")) == isolated_home / "baska"


# --- uctan uca API --------------------------------------------------------


def akisi_kos(api_client, context, fake_gorusme) -> int:
    """Gorusme basla -> bitir -> kuyruk -> not hazir (gercek uygulama uzerinde).

    Asgari sure sifira cekilir: testin gercek saati beklemesi gerekmesin.
    """
    context.settings.set("calls.takip", "1")
    context.settings.set("calls.min_dakika", "0")
    servis = context.gorusme
    fake_gorusme["algilayici"].basla()
    not_id = servis.takipci.yokla()["id"]
    fake_gorusme["algilayici"].bitir()
    servis.takipci.yokla()
    servis.kuyruk.sirayi_isle()
    return not_id


def test_the_whole_flow_runs_through_the_api(api_client, context, fake_gorusme):
    not_id = akisi_kos(api_client, context, fake_gorusme)

    view = api_client.get("/api/gorusme/view").json()
    kart = next(satir for satir in view["notlar"] if satir["id"] == not_id)
    assert kart["durum"] == DURUM_HAZIR
    assert kart["durum_label"] == "hazır"
    assert kart["baslik"] == sahte.ORNEK_OZET["baslik"]
    assert view["serit"]["takip"] is True
    assert view["serit"]["kuyrukta"] == 0

    detay = api_client.get(f"/api/gorusme/{not_id}").json()
    assert [satir["metin"] for satir in detay["bolumler"]["ozet"]] == sahte.ORNEK_OZET["ozet"]
    assert detay["bolum_basliklari"]["aksiyon"] == "Aksiyonlar"
    assert detay["transkript_var"] is False

    # Aksiyondan gorev: Gorevlerim'de gorunur ve nota baglidir.
    aksiyon = detay["bolumler"]["aksiyon"][0]
    uretilen = api_client.post(
        f"/api/gorusme/{not_id}/gorev", json={"bolum_id": aksiyon["id"]}
    ).json()
    assert uretilen["gorev"]["title"] == "Geri dönüş planını yaz"
    pano = api_client.get("/api/tasks").json()
    kartlar = [kart for sutun in pano["columns"] for kart in sutun["tasks"]]
    assert any(kart["title"] == "Geri dönüş planını yaz" for kart in kartlar)

    # Jira bagi: kayit detayindaki "Görüşme notları" bolumu bunu okur.
    baglandi = api_client.put(f"/api/gorusme/{not_id}/jira", json={"jira_key": "PRJ-1432"}).json()
    assert baglandi["jira_key"] == "PRJ-1432"
    kayit = api_client.get("/api/gorusme/kayit/PRJ-1432").json()
    assert kayit["count"] == 1

    silindi = api_client.delete(f"/api/gorusme/{not_id}")
    assert silindi.json()["ok"] is True
    assert api_client.get("/api/gorusme/view").json()["notlar"] == []


def test_the_tracking_switch_can_be_paused_from_the_api(api_client, context):
    context.settings.set("calls.takip", "1")
    serit = api_client.post("/api/gorusme/takip", json={"acik": False}).json()
    assert serit["takip"] is False
    assert context.settings.get("calls.takip") == "0"

    serit = api_client.post("/api/gorusme/takip", json={"acik": True}).json()
    assert serit["takip"] is True


def test_a_failed_note_can_be_retried_from_the_api(api_client, context, fake_gorusme):
    not_id = akisi_kos(api_client, context, fake_gorusme)
    depo.hataya_dus(context.conn, not_id, "yazıya dökülemedi")

    yanit = api_client.post(f"/api/gorusme/{not_id}/yeniden-dene").json()
    assert yanit["not"]["durum"] == DURUM_KUYRUKTA


def test_resummarising_over_the_api_needs_the_transcript(api_client, context, fake_gorusme):
    not_id = akisi_kos(api_client, context, fake_gorusme)
    yanit = api_client.post(f"/api/gorusme/{not_id}/yeniden-ozetle")
    assert yanit.status_code == 400
    assert yanit.json()["error"]["code"] == "transkript_yok"


def test_the_person_drawer_lists_the_notes_of_that_person(api_client, context, fake_gorusme):
    not_id = akisi_kos(api_client, context, fake_gorusme)
    depo.katilimcilari_yaz(context.conn, not_id, [{"kimlik": PERSON_ONE, "ad": "Örnek Kişi"}])

    yanit = api_client.get(f"/api/gorusme/kisi/{PERSON_ONE}").json()
    assert yanit["count"] == 1
    assert yanit["notlar"][0]["katilimcilar"] == ["Örnek Kişi"]


def test_the_settings_endpoints_answer_without_windows(api_client):
    aygitlar = api_client.get("/api/gorusme/ayar/aygitlar").json()
    # Linux'ta ses paketi yok: liste bos, "supported" yanlis, uc yine de calisir.
    assert aygitlar["supported"] is False
    assert aygitlar["aygitlar"] == []
    assert aygitlar["paketler"]["faster_whisper"] in (True, False)

    sablon = api_client.get("/api/gorusme/ayar/sablon").json()
    assert '"aksiyonlar"' in sablon["varsayilan"]
    assert sablon["modeller"][0] == "claude-sonnet-5"


def test_the_trial_endpoint_returns_the_diagnosis_instead_of_a_500(api_client, context):
    """Deneme kaydi asla 500 dondurmez: veri gelmeyen kanal da 200 ve metin."""
    kart = sahte.SahteSesKarti(bloklar=0)
    context.gorusme_kayit_factory = lambda *args, **kwargs: kayitci_kur(kart)
    context.settings.set("calls.mikrofon", "0")
    context.settings.set("calls.hoparlor", "1")

    cevap = api_client.post("/api/gorusme/ayar/deneme", json={"saniye": 0.1})
    assert cevap.status_code == 200
    govde = cevap.json()
    assert govde["hata"] == ""
    hop = govde["kanallar"][KANAL_HOP]
    assert hop["acildi"] is True and hop["cerceve"] == 0
    assert "veri gelmedi" in hop["hata"]
    assert hop["aygit"]["ad"] == "Örnek Hoparlör (döngü)"
    assert govde["oneriler"]


def test_an_unexpected_sound_driver_error_is_shown_as_text(api_client, context):
    def patlayan(*args, **kwargs):
        raise RuntimeError("pyaudio: [Errno -9999] Unanticipated host error\n  ikinci satır")

    context.gorusme_kayit_factory = patlayan
    cevap = api_client.post("/api/gorusme/ayar/deneme", json={"saniye": 0.1})
    assert cevap.status_code == 200
    govde = cevap.json()
    assert "Unanticipated host error" in govde["hata"]
    assert "\n" not in govde["hata"]  # tek satirlik temiz metin
    assert govde["kod"] == "deneme_dustu"
    assert govde["kanallar"][KANAL_MIK]["cerceve"] == 0


def test_the_trial_endpoint_explains_a_missing_capture_package(api_client, context):
    def eksik(*args, **kwargs):
        raise kayit_modulu.eksik_paket()

    context.gorusme_kayit_factory = eksik
    cevap = api_client.post("/api/gorusme/ayar/deneme", json={"saniye": 0.1})
    assert cevap.status_code == 200
    govde = cevap.json()
    assert govde["kod"] == "ses_yakalama_yok"
    assert "pyaudiowpatch" in govde["hata"]


def test_copilot_can_be_tested_from_the_settings_page(api_client, context, fake_gorusme):
    fake_gorusme["ozetleyici"].reddedilen = {"claude-sonnet-5"}
    yanit = api_client.post("/api/gorusme/ayar/copilot-sina").json()
    assert yanit["calisiyor"] is True
    assert yanit["model"] == "gpt-5-mini"
    assert context.settings.get("calls.ozet_model_son") == "gpt-5-mini"
    # Vekil adresi asla geri yazilmaz, yalnizca "ayarli mi" bilgisi doner.
    assert yanit["proxy_ayarli"] is False
    assert "proxy" not in yanit["mesaj"].lower()


def test_the_copilot_path_can_be_tested_and_is_remembered(
    api_client, context, fake_gorusme
):
    """"Copilot'u sına" ekrandaki yolu kullanir, bulunan yolu gosterir ve saklar."""
    yol = r"C:\Users\ornek\AppData\Roaming\npm\copilot.cmd"
    gorulen: dict[str, str] = {}

    def fabrika(ayarlar):
        gorulen["yol"] = ayarlar.copilot_yolu
        fake_gorusme["ozetleyici"].son_yol = yol
        return fake_gorusme["ozetleyici"]

    context.gorusme_ozetleyici_factory = fabrika

    yanit = api_client.post("/api/gorusme/ayar/copilot-sina", json={"yol": yol}).json()

    assert yanit["calisiyor"] is True
    # Alan kaydedilmeden denenebilir: deger uca ekrandan gitti.
    assert gorulen["yol"] == yol
    # Basarida bulunan tam yol ekranda durur ve ayara "son bulunan" yazilir.
    assert yanit["yol"] == yol
    assert yanit["mesaj"].startswith("çalışıyor · " + yol + " · model ")
    assert context.settings.get("calls.copilot_yolu_son") == yol


def test_the_transcription_package_can_be_installed_from_the_settings_page(
    api_client, context, fake_gorusme
):
    not_id = akisi_kos(api_client, context, fake_gorusme)
    depo.hataya_dus(context.conn, not_id, str(yaziyadok.eksik_paket()))
    context.settings.set("calls.copilot_proxy", "http://vekil.ornek.local:8080")

    yanit = api_client.post("/api/gorusme/ayar/whisper-kur").json()

    assert yanit["kuruldu"] is True
    assert yanit["mesaj"]
    # Bekleyen satir kullanicidan "Yeniden dene" beklemeden kuyruga dondu.
    # (Arka plan iscisi onu hemen alip yeniden isleyebilir; kesin olan sey
    # satirin artik "paket yok" hatasinda BEKLEMEDIGI.)
    assert yanit["yeniden_kuyruga"] == 1
    assert "faster-whisper" not in depo.require_not(context.conn, not_id)["hata"]
    # Vekil kurucuya verildi; Holocron'un kendi ortami degismedi.
    assert fake_gorusme["kurucu"].cagrilar == ["http://vekil.ornek.local:8080"]
    assert "HTTPS_PROXY" not in os.environ


def test_a_failed_installation_leaves_the_queue_alone(api_client, context, fake_gorusme):
    not_id = akisi_kos(api_client, context, fake_gorusme)
    depo.hataya_dus(context.conn, not_id, str(yaziyadok.eksik_paket()))
    fake_gorusme["kurucu"].kuruldu = False

    yanit = api_client.post("/api/gorusme/ayar/whisper-kur").json()

    assert yanit["kuruldu"] is False
    assert yanit["yeniden_kuyruga"] == 0
    assert depo.require_not(context.conn, not_id)["durum"] == DURUM_HATA


def test_a_note_never_leaks_raw_identities_to_the_screen(api_client, context, fake_gorusme):
    not_id = akisi_kos(api_client, context, fake_gorusme)
    depo.katilimcilari_yaz(context.conn, not_id, [{"kimlik": PERSON_ONE, "ad": ""}])
    kart = api_client.get("/api/gorusme/view").json()["notlar"][0]
    assert PERSON_ONE not in " ".join(kart["katilimcilar"])
    assert kart["katilimcilar"][0].startswith("Bilinmeyen kişi")


# --- arayuz ---------------------------------------------------------------


def test_the_notes_tab_and_the_track_strip_are_on_the_page(api_client):
    page = api_client.get("/").text
    for marker in (
        'id="calls-tab-notes"',
        "Görüşme notları",
        'id="gorusme-track"',
        'id="gorusme-takip"',
        'id="gorusme-takip-toggle"',
        'id="gorusme-live"',
        'id="gorusme-working"',
        'id="gorusme-queued"',
        'id="gorusme-count"',
        'id="gorusme-wrap"',
        'id="gorusme-whisper-uyari"',
        "Yazıya dökme paketi eksik",
        "Ayarlar'dan kur",
        'id="gorusme-whisper-uyari-kapat"',
        'id="gorusme-table"',
        'id="gorusme-drawer"',
        "/static/js/gorusme.js",
    ):
        assert marker in page, marker
    # Sekme Gruplar'dan sonra gelir.
    assert page.index('id="calls-tab-groups"') < page.index('id="calls-tab-notes"')


def test_the_notes_script_covers_the_strip_the_table_and_the_drawer(api_client):
    script = api_client.get("/static/js/gorusme.js").text
    for marker in (
        "/api/gorusme/view",
        "/api/gorusme/serit",
        "/api/gorusme/takip",
        "/api/gorusme/kayit/",
        "/api/gorusme/kisi/",
        "function renderGorusmeTrack",
        "function renderGorusmeList",
        "function openGorusme",
        "function makeGorusmeTask",
        "function bindGorusmeJira",
        "function retryGorusme",
        "function resummarizeGorusme",
        "function renderDrawerGorusme",
        "function renderGorusmeWhisperWarning",
        "function dismissWhisperWarning",
        "function renderKisiGorusme",
        '"Görev yap"',
        '"Yeniden dene"',
        '"Yeniden özetle"',
        '"E-posta ile gönder"',
        "Duraklat",
        "Sürdür",
    ):
        assert marker in script, marker
    # Türkçe metne büyük harf dönüşümü uygulanmaz.
    assert "toUpperCase()" not in script


def test_the_calls_screen_hands_the_fourth_tab_over(api_client):
    script = api_client.get("/static/js/calls.js").text
    assert 'notes: "calls-tab-notes"' in script
    assert "showGorusmeTab" in script
    assert "renderKisiGorusme" in script


def test_the_note_styles_are_defined(api_client):
    css = api_client.get("/static/css/app.css").text
    for name in (
        ".track-strip",
        ".pill.live",
        ".pill.work",
        ".pill.queue",
        ".pill.state-hazir",
        ".pill.state-hata",
        ".bar",
        ".gorusme-row",
        ".who-chips",
        ".action-row",
        ".jira-box",
        ".note-list",
    ):
        assert name in css, name
    # Turkce metinde buyuk harf donusumu yasak.
    assert "text-transform: uppercase" not in css.split(".track-strip")[1][:2000]


def test_the_settings_card_is_on_the_settings_page(api_client):
    page = api_client.get("/settings").text
    for marker in (
        'id="gorusme-card"',
        "Görüşme notları",
        'id="gorusme-mikrofon"',
        'id="gorusme-hoparlor"',
        'id="gorusme-deneme"',
        "Deneme kaydı (10 sn)",
        'id="gorusme-min-dakika"',
        'id="gorusme-whisper-model"',
        'id="gorusme-whisper-klasor"',
        'id="gorusme-whisper-kur"',
        "Yazıya dökme paketini kur",
        'id="gorusme-whisper-kur-sonuc"',
        'id="gorusme-model-sina"',
        "Modeli sına",
        'id="gorusme-model-sina-sonuc"',
        "model.bin",
        'id="gorusme-modeller"',
        'id="gorusme-copilot-proxy"',
        'id="gorusme-copilot-yol"',
        'id="gorusme-copilot-sonuc"',
        'id="gorusme-copilot-test"',
        "Copilot'u sına",
        "where copilot",
        'id="gorusme-sablon"',
        'id="gorusme-klasor"',
        'id="gorusme-transkript-sakla"',
        'id="gorusme-bildirim-hazir"',
        'id="gorusme-isleme-disinda"',
        "/static/js/gorusme-settings.js",
    ):
        assert marker in page, marker
    # Vekil alani ne yaptigini anlatir: Jira dogrudan kalir.
    assert "doğrudan bağlan" in page


def test_the_settings_script_saves_every_field(api_client):
    script = api_client.get("/static/js/gorusme-settings.js").text
    for marker in (
        "/api/gorusme/ayar/aygitlar",
        "/api/gorusme/ayar/deneme",
        "/api/gorusme/ayar/copilot-sina",
        "/api/gorusme/ayar/sablon",
        '"calls.copilot_proxy"',
        '"calls.copilot_yolu"',
        '"calls.ozet_modelleri"',
        '"calls.min_dakika"',
        '"calls.transkripti_sakla"',
        '"calls.isleme_gorusme_disinda"',
        "ses var",
        "ses yok",
        # Kanal basina teshis satiri: aygit adi, acildi mi, kac cerceve geldi.
        "çerçeve",
        "açılamadı",
        "sonuc.oneriler",
        "sonuc.hata",
        # Uc `{settings: {...}}` doner: sozlugu acmadan okumak butun onay
        # kutularini bos gosteriyordu.
        "data.settings",
    ):
        assert marker in script, marker


def test_the_export_carries_a_meeting_notes_sheet(api_client, context, fake_gorusme):
    from openpyxl import load_workbook
    import io as stdio

    not_id = akisi_kos(api_client, context, fake_gorusme)
    depo.jira_bagla(context.conn, not_id, "PRJ-1432")

    yanit = api_client.get("/api/calls/export.xlsx?days=90")
    book = load_workbook(stdio.BytesIO(yanit.content))
    assert "Görüşme notları" in book.sheetnames

    sheet = book["Görüşme notları"]
    basliklar = [cell.value for cell in sheet[1]]
    assert basliklar[:3] == ["Tarih", "Başlık", "Katılımcılar"]
    assert "Aksiyonlar" in basliklar
    assert sheet.cell(row=2, column=2).value == sahte.ORNEK_OZET["baslik"]
    assert sheet.cell(row=2, column=8).value == "PRJ-1432"
    # Özet maddeleri hücre içinde satır satır durur.
    assert sahte.ORNEK_OZET["ozet"][0] in str(sheet.cell(row=2, column=9).value)

"""Grup -> Excel disa aktarimi: uretilen dosya openpyxl ile geri okunur."""

from __future__ import annotations

import io
from datetime import date, datetime

from openpyxl import load_workbook

from app import export
from app import repository as repo
from app.settings_store import MODE_SERVER
from tests.fake_jira import issue

BASE = "https://jira.example.com"

CATALOG = [
    {"id": "summary", "name": "Özet", "schema": {"type": "string"}},
    {"id": "status", "name": "Durum", "schema": {"type": "status"}},
    {"id": "duedate", "name": "Bitiş tarihi", "schema": {"type": "date"}},
    {"id": "updated", "name": "Güncellendi", "schema": {"type": "datetime"}},
    {"id": "storypoints", "name": "Puan", "custom": True, "schema": {"type": "number"}},
]

# Grid 200 karakterde kirpar; Excel'de tam metin bekleriz.
LONG_TEXT = "Beta " + ("çok uzun bir özet " * 20)

COLUMNS = ["issuekey", "summary", "status", "duedate", "updated", "storypoints"]


def setup_group(api_client, conn, columns=None):
    repo.store_fields(conn, CATALOG)
    api_client.put("/api/settings", json={"jira.base_url": BASE, "jira.mode": MODE_SERVER})
    group = api_client.post("/api/groups", json={"name": "Filom", "kind": "manual"}).json()["group"]
    api_client.post(f"/api/groups/{group['id']}/items", json={"text": "DEMO-1 DEMO-2"})
    repo.upsert_issues(
        conn,
        [
            issue(
                "DEMO-1",
                "Alfa kaydı",
                status={"name": "Açık"},
                duedate="2026-09-11",
                updated="2026-09-11T09:05:00+0300",
                storypoints=3.5,
            ),
            issue(
                "DEMO-2",
                LONG_TEXT,
                status={"name": "Kapalı"},
                duedate=None,
                updated="2026-01-02T23:45:00+0000",
                storypoints=12,
            ),
        ],
    )
    choose(api_client, group, columns or COLUMNS)
    return group


def choose(api_client, group, columns):
    response = api_client.put(f"/api/groups/{group['id']}", json={"columns": columns})
    assert response.status_code == 200, response.text


def add_field(api_client, name, type="text", options=None, history=False):
    payload = {"name": name, "type": type, "track_history": history}
    if options is not None:
        payload["options"] = options
    response = api_client.post("/api/local-fields", json=payload)
    assert response.status_code == 200, response.text
    return response.json()["field"]


def download(api_client, group, **params):
    response = api_client.get(f"/api/groups/{group['id']}/export.xlsx", params=params)
    assert response.status_code == 200, response.text
    return response


def opened(response):
    return load_workbook(io.BytesIO(response.content))


def values(sheet, row):
    return [cell.value for cell in sheet[row]]


def local_moment(text: str) -> datetime:
    """Beklenen Excel degeri: kaydin zamani makinenin yerel saatinde."""
    return datetime.fromisoformat(text).astimezone().replace(tzinfo=None)


# --- sayfa 1: kayitlar --------------------------------------------------


def test_headers_use_human_names_in_column_order(api_client, conn):
    group = setup_group(api_client, conn)
    sheet = opened(download(api_client, group)).worksheets[0]

    assert sheet.title == "Filom"
    assert values(sheet, 1) == [
        "Anahtar",
        "Özet",
        "Durum",
        "Bitiş tarihi",
        "Güncellendi",
        "Puan",
    ]
    assert all(cell.font.bold for cell in sheet[1])
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:F3"
    assert [sheet.cell(row=line, column=1).value for line in (2, 3)] == ["DEMO-1", "DEMO-2"]


def test_date_and_datetime_cells_are_real_dates(api_client, conn):
    group = setup_group(api_client, conn)
    sheet = opened(download(api_client, group)).worksheets[0]

    due = sheet.cell(row=2, column=4)
    assert isinstance(due.value, datetime)
    assert due.value.date() == date(2026, 9, 11)
    assert due.number_format == "DD.MM.YYYY"

    updated = sheet.cell(row=2, column=5)
    assert isinstance(updated.value, datetime)
    assert updated.value == local_moment("2026-09-11T09:05:00+03:00")
    assert updated.number_format == "DD.MM.YYYY HH:MM"

    # Bos tarih metne dusmez, hucre bos kalir.
    assert sheet.cell(row=3, column=4).value is None


def test_number_cells_keep_their_type(api_client, conn):
    group = setup_group(api_client, conn)
    sheet = opened(download(api_client, group)).worksheets[0]

    assert sheet.cell(row=2, column=6).value == 3.5
    assert isinstance(sheet.cell(row=2, column=6).value, float)
    assert sheet.cell(row=3, column=6).value == 12
    assert isinstance(sheet.cell(row=3, column=6).value, int)


def test_long_text_is_not_trimmed(api_client, conn):
    group = setup_group(api_client, conn)
    sheet = opened(download(api_client, group)).worksheets[0]

    grid = api_client.get(f"/api/groups/{group['id']}/issues").json()
    shown = [row for row in grid["rows"] if row["key"] == "DEMO-2"][0]["cells"][1]["text"]

    assert shown.endswith("…") and len(shown) == 200
    assert sheet.cell(row=3, column=2).value == LONG_TEXT.strip()
    assert len(sheet.cell(row=3, column=2).value) > 200


def test_key_cell_links_to_jira(api_client, conn):
    group = setup_group(api_client, conn)
    sheet = opened(download(api_client, group)).worksheets[0]

    cell = sheet.cell(row=2, column=1)
    assert cell.value == "DEMO-1"
    assert cell.hyperlink.target == f"{BASE}/browse/DEMO-1"


def test_without_base_url_key_stays_plain(api_client, conn):
    group = setup_group(api_client, conn)
    api_client.put("/api/settings", json={"jira.base_url": ""})
    sheet = opened(download(api_client, group)).worksheets[0]

    assert sheet.cell(row=2, column=1).value == "DEMO-1"
    assert sheet.cell(row=2, column=1).hyperlink is None


def test_column_choice_is_respected(api_client, conn):
    group = setup_group(api_client, conn)
    sheet = opened(download(api_client, group, columns="issuekey,storypoints")).worksheets[0]

    assert values(sheet, 1) == ["Anahtar", "Puan"]
    assert sheet.max_column == 2


def test_query_and_sort_are_applied(api_client, conn):
    group = setup_group(api_client, conn)

    filtered = opened(download(api_client, group, q="alfa")).worksheets[0]
    assert [filtered.cell(row=line, column=1).value for line in (2, 3)] == ["DEMO-1", None]

    sorted_sheet = opened(
        download(api_client, group, sort="storypoints", dir="desc")
    ).worksheets[0]
    assert [sorted_sheet.cell(row=line, column=1).value for line in (2, 3)] == ["DEMO-2", "DEMO-1"]


def test_local_and_derived_cells_carry_their_type(api_client, conn):
    group = setup_group(api_client, conn)
    note = add_field(api_client, "Puanım", type="number")
    when = add_field(api_client, "Söz verilen", type="date")
    flag = add_field(api_client, "Faturalı", type="bool")
    status = add_field(api_client, "Müşteri durumu", type="select", options=["bekliyor", "bitti"], history=True)
    base = status["column_id"]
    choose(
        api_client,
        group,
        [
            "issuekey",
            note["column_id"],
            when["column_id"],
            flag["column_id"],
            base,
            f"{base}:changed_at",
            f"{base}:changes",
        ],
    )
    api_client.put(f"/api/issues/DEMO-1/local/{note['id']}", json={"value": "12,50"})
    api_client.put(f"/api/issues/DEMO-1/local/{when['id']}", json={"value": "11.09.2026"})
    api_client.put(f"/api/issues/DEMO-1/local/{flag['id']}", json={"value": "evet"})
    api_client.put(f"/api/issues/DEMO-2/local/{flag['id']}", json={"value": "hayir"})
    api_client.put(f"/api/issues/DEMO-1/local/{status['id']}", json={"value": "bekliyor"})
    api_client.put(f"/api/issues/DEMO-1/local/{status['id']}", json={"value": "bitti"})

    sheet = opened(download(api_client, group)).worksheets[0]
    assert values(sheet, 1) == [
        "Anahtar",
        "Puanım",
        "Söz verilen",
        "Faturalı",
        "Müşteri durumu",
        "Müşteri durumu (son değişim)",
        "Müşteri durumu (kaç kez değişti)",
    ]
    assert sheet.cell(row=2, column=2).value == 12.5
    assert sheet.cell(row=2, column=3).value == datetime(2026, 9, 11)
    assert sheet.cell(row=2, column=3).number_format == "DD.MM.YYYY"
    assert sheet.cell(row=2, column=4).value == "Evet"
    assert sheet.cell(row=3, column=4).value == "Hayır"
    assert sheet.cell(row=2, column=5).value == "bitti"
    assert isinstance(sheet.cell(row=2, column=6).value, datetime)
    assert sheet.cell(row=2, column=7).value == 2
    assert sheet.cell(row=3, column=7).value == 0


# --- sayfa 2: gecmis ----------------------------------------------------


def test_history_sheet_lists_changes_newest_first(api_client, conn):
    group = setup_group(api_client, conn)
    status = add_field(
        api_client, "Müşteri durumu", type="select", options=["bekliyor", "bitti"], history=True
    )
    note = add_field(api_client, "Not", type="text", history=False)
    choose(api_client, group, ["issuekey", status["column_id"], note["column_id"]])
    api_client.put(f"/api/issues/DEMO-1/local/{status['id']}", json={"value": "bekliyor"})
    api_client.put(f"/api/issues/DEMO-1/local/{status['id']}", json={"value": "bitti"})
    api_client.put(f"/api/issues/DEMO-2/local/{note['id']}", json={"value": "gecmisi yok"})

    book = opened(download(api_client, group, history="1"))
    assert book.sheetnames == ["Filom", "Geçmiş", "Bilgi"]

    sheet = book["Geçmiş"]
    assert values(sheet, 1) == ["Anahtar", "Alan", "Tarih", "Eski değer", "Yeni değer"]
    assert sheet.max_row == 3
    assert [sheet.cell(row=2, column=index).value for index in (1, 2, 4, 5)] == [
        "DEMO-1",
        "Müşteri durumu",
        "bekliyor",
        "bitti",
    ]
    assert isinstance(sheet.cell(row=2, column=3).value, datetime)
    # Ilk degisimde eski deger yok; gecmisi kapali alan sayfaya girmez.
    assert sheet.cell(row=3, column=4).value is None
    assert sheet.cell(row=3, column=5).value == "bekliyor"


def test_history_sheet_skips_records_outside_the_filter(api_client, conn):
    group = setup_group(api_client, conn)
    status = add_field(
        api_client, "Durumum", type="select", options=["bekliyor", "bitti"], history=True
    )
    choose(api_client, group, ["issuekey", "summary", status["column_id"]])
    api_client.put(f"/api/issues/DEMO-1/local/{status['id']}", json={"value": "bekliyor"})
    api_client.put(f"/api/issues/DEMO-2/local/{status['id']}", json={"value": "bitti"})

    sheet = opened(download(api_client, group, history="1", q="alfa"))["Geçmiş"]
    assert sheet.max_row == 2
    assert sheet.cell(row=2, column=1).value == "DEMO-1"


def test_history_sheet_without_tracked_fields_says_so(api_client, conn):
    group = setup_group(api_client, conn)
    add_field(api_client, "Not", type="text", history=False)

    sheet = opened(download(api_client, group, history="1"))["Geçmiş"]
    assert sheet["A1"].value == "Geçmiş tutulan alan yok"
    assert sheet.max_row == 1


def test_history_sheet_is_absent_unless_asked(api_client, conn):
    group = setup_group(api_client, conn)
    assert opened(download(api_client, group)).sheetnames == ["Filom", "Bilgi"]


# --- sayfa 3: bilgi -----------------------------------------------------


def test_info_sheet_describes_the_export(api_client, conn):
    group = setup_group(api_client, conn)
    sheet = opened(download(api_client, group))["Bilgi"]
    info = {row[0].value: row[1].value for row in sheet.iter_rows(min_row=1, max_col=2)}

    assert info["Grup"] == "Filom"
    assert info["Tür"] == "Manuel"
    assert "JQL" not in info
    assert info["Kayıt sayısı"] == 2
    assert isinstance(info["Dışa aktarma"], datetime)
    from app import __version__

    assert info["Sürüm"] == __version__


def test_info_sheet_carries_jql_for_filter_groups(api_client, conn):
    repo.store_fields(conn, CATALOG)
    group = api_client.post(
        "/api/groups", json={"name": "Sorgu", "kind": "filter", "jql": "project = DEMO"}
    ).json()["group"]

    sheet = opened(download(api_client, group))["Bilgi"]
    info = {row[0].value: row[1].value for row in sheet.iter_rows(min_row=1, max_col=2)}
    assert info["Tür"] == "JQL filtresi"
    assert info["JQL"] == "project = DEMO"
    assert info["Kayıt sayısı"] == 0


# --- adlandirma ve uc davranisi -----------------------------------------


def test_sheet_title_is_cleaned_and_shortened():
    assert export.sheet_title("Filo: [Müşteri]/2026") == "Filo Müşteri 2026"
    assert len(export.sheet_title("A" * 60)) == 31
    assert export.sheet_title("   ") == "Kayıtlar"
    assert export.sheet_title("'alinti'") == "alinti"


def test_group_name_becomes_the_sheet_title(api_client, conn):
    group = setup_group(api_client, conn)
    api_client.put(f"/api/groups/{group['id']}", json={"name": "Müşteri: A/B [2026]"})
    book = opened(download(api_client, group))
    assert book.worksheets[0].title == "Müşteri A B 2026"


def test_content_disposition_carries_utf8_and_ascii_names(api_client, conn):
    group = setup_group(api_client, conn)
    api_client.put(f"/api/groups/{group['id']}", json={"name": "Müşteri İşleri"})
    response = download(api_client, group)

    stamp = date.today().isoformat()
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment; ")
    assert f'filename="Musteri-Isleri-{stamp}.xlsx"' in disposition
    assert f"filename*=UTF-8''M%C3%BC%C5%9Fteri-%C4%B0%C5%9Fleri-{stamp}.xlsx" in disposition
    assert response.headers["content-type"] == export.MEDIA_TYPE


def test_empty_group_still_produces_a_workbook(api_client, conn):
    group = api_client.post("/api/groups", json={"name": "Boş"}).json()["group"]
    book = opened(download(api_client, group, history="1"))

    sheet = book.worksheets[0]
    assert sheet.max_row == 1
    assert sheet.cell(row=1, column=1).value == "Anahtar"
    assert book["Bilgi"]["B4"].value == 0


def test_unknown_group_is_404_in_common_shape(api_client):
    response = api_client.get("/api/groups/999/export.xlsx")
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "group_not_found"

"""Asama 8: kayittan Teams'e mesaj.

Uc katman ayri ayri sinanir: saf baglanti/sablon mantigi (`app/teams.py`),
depo ve uclar (kisi, sablon, rehber, "acildi" kaydi), son olarak sanal sutunun
grid ve Excel'e dusmesi. Gercek Teams'e cikilmaz: uretilen tek sey bir adres.
"""

from __future__ import annotations

import io
from urllib.parse import parse_qs, unquote, urlparse

import pytest
from openpyxl import load_workbook

from app import db
from app import desktop
from app import repository as repo
from app import teams
from app.mail import GalEntry, MailError
from app.mail import outlook
from app.mail.source import unavailable as mail_unavailable
from app.settings_store import MODE_SERVER
from tests.fake_jira import issue

BASE = "https://jira.example.com"

CATALOG = [
    {"id": "summary", "name": "Özet", "schema": {"type": "string"}},
    {"id": "status", "name": "Durum", "schema": {"type": "status"}},
    {"id": "assignee", "name": "Atanan", "schema": {"type": "user"}},
    {"id": "priority", "name": "Öncelik", "schema": {"type": "priority"}},
    {"id": "customfield_10016", "name": "Puan", "custom": True, "schema": {"type": "number"}},
]

ORNEK = "ornek@example.com"
IKINCI = "ikinci@example.com"


def setup_group(api_client, conn, keys=("DEMO-1",)):
    repo.store_fields(conn, CATALOG)
    api_client.put("/api/settings", json={"jira.base_url": BASE, "jira.mode": MODE_SERVER})
    group = api_client.post("/api/groups", json={"name": "Filom", "kind": "manual"}).json()["group"]
    api_client.post(f"/api/groups/{group['id']}/items", json={"text": " ".join(keys)})
    repo.upsert_issues(
        conn,
        [
            issue(
                key,
                "Yazıcı arızası",
                status={"name": "Beklemede", "statusCategory": {"key": "indeterminate"}},
                assignee={"displayName": "Ali Veli"},
                priority={"name": "Yüksek"},
                customfield_10016=5,
            )
            for key in keys
        ],
    )
    return group


def add_contacts(api_client, key="DEMO-1", contacts=None):
    payload = contacts if contacts is not None else [{"email": ORNEK, "name": "Örnek Kişi"}]
    response = api_client.put(f"/api/issues/{key}/contacts", json={"contacts": payload})
    assert response.status_code == 200, response.text
    return response.json()["contacts"]


def query_of(url):
    return parse_qs(urlparse(url).query, keep_blank_values=True)


# --- goc 0006 -----------------------------------------------------------


def test_migration_creates_the_teams_tables_and_seeds_templates(conn):
    tables = db.table_names(conn)
    for name in ("contacts", "issue_contacts", "message_templates", "sent_messages"):
        assert name in tables, name
    assert db.current_version(conn) == 7

    seeded = repo.list_templates(conn)
    assert [item["name"] for item in seeded] == ["Son durum", "Güncelleme rica"]
    assert seeded[0]["is_default"] is True
    assert "{key}" in seeded[0]["body"] and "{status}" in seeded[0]["body"]


def test_migration_does_not_bring_back_deleted_templates(conn):
    for template in repo.list_templates(conn):
        repo.delete_template(conn, template["id"])
    # Yarim kalmis bir yukseltme tekrar calisirsa tohum atilmamali.
    db._migration_0006_teams(conn)
    assert repo.list_templates(conn) == []


# --- saf mantik: sablon cozumu ------------------------------------------


def test_render_template_fills_every_placeholder(conn):
    record = repo.get_issue(conn, "DEMO-1")
    raw = issue(
        "DEMO-1",
        "Yazıcı arızası",
        status={"name": "Beklemede"},
        assignee={"displayName": "Ali Veli"},
        priority={"name": "Yüksek"},
        customfield_10016=5,
    )
    body = (
        "{key} / {summary} / {status} / {assignee} / {priority} / {url} / "
        "{field:customfield_10016} / {local:3} / {yok}"
    )
    text = teams.render_template(body, raw, {3: "müşteriye soruldu"}, BASE)
    assert text == (
        "DEMO-1 / Yazıcı arızası / Beklemede / Ali Veli / Yüksek / "
        f"{BASE}/browse/DEMO-1 / 5 / müşteriye soruldu / "
    )
    assert record is None  # bu test depoya hic dokunmaz


def test_render_template_accepts_the_stored_record_shape():
    stored = {"key": "DEMO-7", "fetched_at": "x", "raw": {"fields": {"summary": "Konu"}}}
    assert teams.render_template("{key}: {summary}", stored, {}, BASE) == "DEMO-7: Konu"


def test_render_template_leaves_unknown_placeholders_empty():
    assert teams.render_template("[{musteri}]", {"key": "DEMO-1"}) == "[]"
    assert teams.render_template("[{local:99}]", {"key": "DEMO-1"}, {3: "var"}) == "[]"
    assert teams.render_template("[{field:yok}]", {"key": "DEMO-1"}) == "[]"


def test_render_template_without_base_url_leaves_the_link_empty():
    assert teams.render_template("{url}", {"key": "DEMO-1"}, {}, "") == ""


def test_local_values_can_be_addressed_with_string_keys():
    assert teams.render_template("{local:3}", {"key": "K-1"}, {"3": "metin"}) == "metin"


# --- saf mantik: baglanti kurma -----------------------------------------


def test_single_person_link_has_no_topic():
    link = teams.build_chat_link([ORNEK], "Merhaba")
    assert link["url"].startswith("https://teams.microsoft.com/l/chat/0/0?users=")
    query = query_of(link["url"])
    assert query["users"] == [ORNEK]
    assert query["message"] == ["Merhaba"]
    assert "topicName" not in query
    assert link["truncated"] is False


def test_group_link_joins_addresses_and_names_the_chat():
    link = teams.build_chat_link([ORNEK, IKINCI], "Durum?", topic="DEMO-1")
    query = query_of(link["url"])
    assert query["users"] == [f"{ORNEK},{IKINCI}"]
    assert query["topicName"] == ["DEMO-1"]
    # Virgul okunakli kalir, mesaj tam kodlanir.
    assert f"users={ORNEK},{IKINCI}" in link["url"]


def test_message_is_percent_encoded_not_plus_encoded():
    link = teams.build_chat_link([ORNEK], "iki kelime & soru?")
    assert "%20" in link["url"] and "+" not in link["url"].split("message=")[1]
    assert unquote(link["url"].split("message=")[1]) == "iki kelime & soru?"


def test_long_message_is_trimmed_to_stay_under_the_url_limit():
    link = teams.build_chat_link([ORNEK], "ç" * 4000)
    assert link["truncated"] is True
    assert len(link["url"]) <= teams.URL_LIMIT
    assert link["message"].endswith(teams.ELLIPSIS)
    assert len(link["message"]) < 4000


def test_short_message_is_not_trimmed():
    link = teams.build_chat_link([ORNEK], "kısa")
    assert link["truncated"] is False and link["message"] == "kısa"


def test_when_the_addresses_alone_exceed_the_limit_the_message_drops():
    crowd = [f"kisi{index}@example.com" for index in range(120)]
    link = teams.build_chat_link(crowd, "Durum?", topic="DEMO-1")
    assert link["truncated"] is True
    assert link["message"] == ""
    assert "message=" not in link["url"]


def test_blank_addresses_are_dropped():
    link = teams.build_chat_link([" ", ORNEK, ""], "Selam")
    assert query_of(link["url"])["users"] == [ORNEK]


def test_the_two_bases_are_the_web_and_the_app_forms():
    # Yedek yol tarayicidan gecer, asil yol isletim sisteminden.
    assert teams.CHAT_BASE.startswith("https://teams.microsoft.com/")
    assert teams.APP_BASE.startswith("msteams:")


# --- kisiler ve adres defteri -------------------------------------------


def test_contacts_round_trip_and_feed_the_address_book(api_client, conn):
    setup_group(api_client, conn)
    saved = add_contacts(
        api_client, contacts=[{"email": "Ornek@Example.COM", "name": "Örnek Kişi"}]
    )
    assert saved == [{"email": ORNEK, "name": "Örnek Kişi", "kind": "person"}]

    read = api_client.get("/api/issues/demo-1/contacts").json()["contacts"]
    assert read == saved
    # Deftere de dusmus olmali; elle girildigi icin kaynagi "manual".
    assert api_client.get("/api/contacts").json()["contacts"] == [
        {"email": ORNEK, "name": "Örnek Kişi", "source": "manual", "kind": "person"}
    ]


def test_put_replaces_the_whole_list(api_client, conn):
    setup_group(api_client, conn)
    add_contacts(api_client, contacts=[ORNEK, IKINCI])
    left = add_contacts(api_client, contacts=[IKINCI])
    assert [item["email"] for item in left] == [IKINCI]
    # Defterden dusmez: kisi baska kayitta lazim olabilir.
    assert len(api_client.get("/api/contacts").json()["contacts"]) == 2


def test_contact_order_is_kept(api_client, conn):
    setup_group(api_client, conn)
    saved = add_contacts(api_client, contacts=[IKINCI, ORNEK])
    assert [item["email"] for item in saved] == [IKINCI, ORNEK]


def test_invalid_email_is_refused(api_client, conn):
    setup_group(api_client, conn)
    response = api_client.put("/api/issues/DEMO-1/contacts", json={"contacts": ["ornek"]})
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "invalid_email"


def test_contacts_on_an_unknown_issue_are_404(api_client, conn):
    setup_group(api_client, conn)
    response = api_client.put("/api/issues/DEMO-404/contacts", json={"contacts": [ORNEK]})
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "issue_not_found"


def test_address_book_search_puts_prefix_matches_first(api_client, conn):
    setup_group(api_client, conn)
    add_contacts(
        api_client,
        contacts=[
            {"email": "veli.ali@example.com", "name": "Veli Ali"},
            {"email": "ali.can@example.com", "name": "Ali Can"},
        ],
    )
    found = api_client.get("/api/contacts", params={"q": "ali"}).json()["contacts"]
    # "Ali Can" onek eslesmesi; "Veli Ali" yalnizca iceriyor.
    assert [item["name"] for item in found] == ["Ali Can", "Veli Ali"]


def test_address_book_search_folds_turkish_letters(api_client, conn):
    setup_group(api_client, conn)
    add_contacts(api_client, contacts=[{"email": "isil@example.com", "name": "Işıl Yıldız"}])
    for needle in ("ışı", "isi", "IŞI", "Isi"):
        found = api_client.get("/api/contacts", params={"q": needle}).json()["contacts"]
        assert [item["email"] for item in found] == ["isil@example.com"], needle


def test_address_book_search_matches_the_email_too(api_client, conn):
    setup_group(api_client, conn)
    add_contacts(api_client, contacts=[{"email": "tedarik@example.com", "name": "Satın alma"}])
    found = api_client.get("/api/contacts", params={"q": "tedarik"}).json()["contacts"]
    assert [item["email"] for item in found] == ["tedarik@example.com"]


def test_address_book_returns_at_most_twenty(api_client, conn):
    setup_group(api_client, conn)
    add_contacts(api_client, contacts=[f"kisi{index}@example.com" for index in range(30)])
    assert len(api_client.get("/api/contacts").json()["contacts"]) == 20


def test_contact_can_be_renamed_and_the_issue_follows(api_client, conn):
    setup_group(api_client, conn)
    add_contacts(api_client)
    response = api_client.put(
        f"/api/contacts/{ORNEK}", json={"email": "yeni@example.com", "name": "Yeni Ad"}
    )
    assert response.status_code == 200, response.text
    assert response.json()["contact"] == {
        "email": "yeni@example.com",
        "name": "Yeni Ad",
        "source": "manual",
        "kind": "person",
    }
    assert api_client.get("/api/issues/DEMO-1/contacts").json()["contacts"] == [
        {"email": "yeni@example.com", "name": "Yeni Ad", "kind": "person"}
    ]


def test_deleting_a_contact_removes_it_from_the_issues(api_client, conn):
    setup_group(api_client, conn)
    add_contacts(api_client, contacts=[ORNEK, IKINCI])
    response = api_client.delete(f"/api/contacts/{ORNEK}")
    assert response.json() == {"ok": True, "issues": 1}
    assert [item["email"] for item in api_client.get("/api/issues/DEMO-1/contacts").json()["contacts"]] == [
        IKINCI
    ]

    missing = api_client.delete(f"/api/contacts/{ORNEK}")
    assert missing.status_code == 404 and missing.json()["error"]["code"] == "contact_not_found"


# --- sablonlar ----------------------------------------------------------


def test_template_crud_and_single_default(api_client):
    listing = api_client.get("/api/templates").json()
    assert listing["topic_format"] == "{key}"
    assert any(item["token"] == "{field:customfield_10016}" for item in listing["placeholders"])
    first = listing["templates"][0]

    created = api_client.post(
        "/api/templates", json={"name": "Kapanış", "body": "{key} kapandı mı?"}
    ).json()["template"]
    assert created["is_default"] is False

    api_client.put(f"/api/templates/{created['id']}", json={"is_default": True})
    after = api_client.get("/api/templates").json()["templates"]
    marked = [item["id"] for item in after if item["is_default"]]
    assert marked == [created["id"]]

    assert api_client.delete(f"/api/templates/{created['id']}").json()["ok"] is True
    # Varsayilan silinince isaret listenin ilkine gecer.
    assert api_client.get("/api/templates").json()["templates"][0]["id"] == first["id"]
    assert api_client.get("/api/templates").json()["templates"][0]["is_default"] is True


def test_empty_template_body_is_refused(api_client):
    response = api_client.post("/api/templates", json={"name": "Boş", "body": "  "})
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "invalid_body"


def test_template_order_can_be_changed(api_client):
    templates = api_client.get("/api/templates").json()["templates"]
    api_client.put(f"/api/templates/{templates[1]['id']}", json={"position": -1})
    assert [item["name"] for item in api_client.get("/api/templates").json()["templates"]] == [
        "Güncelleme rica",
        "Son durum",
    ]


# --- baglanti ucu -------------------------------------------------------


def test_teams_link_uses_the_default_template_and_records_the_open(api_client, conn):
    setup_group(api_client, conn)
    add_contacts(api_client)

    data = api_client.post("/api/issues/DEMO-1/teams-link", json={}).json()
    assert data["kind"] == "people"
    assert data["truncated"] is False
    assert data["message"] == (
        "Merhaba, DEMO-1 (Yazıcı arızası) kaydının son durumu nedir? "
        f"Durum: Beklemede. {BASE}/browse/DEMO-1"
    )
    assert query_of(data["url"])["users"] == [ORNEK]

    history = api_client.get("/api/issues/DEMO-1/messages").json()["messages"]
    assert len(history) == 1
    assert history[0]["target_kind"] == "people"
    assert history[0]["target_text"] == ORNEK
    assert history[0]["first_line"].startswith("Merhaba, DEMO-1")
    assert history[0]["opened_at"]


def test_teams_link_names_the_group_chat_with_the_topic_format(api_client, conn):
    setup_group(api_client, conn)
    add_contacts(api_client, contacts=[ORNEK, IKINCI])
    api_client.put("/api/settings", json={"teams.topic_format": "{key} · {summary}"})

    data = api_client.post("/api/issues/DEMO-1/teams-link", json={}).json()
    assert query_of(data["url"])["topicName"] == ["DEMO-1 · Yazıcı arızası"]


def test_teams_link_accepts_a_hand_written_body(api_client, conn):
    setup_group(api_client, conn)
    add_contacts(api_client)
    data = api_client.post(
        "/api/issues/DEMO-1/teams-link", json={"body": "{key} ne durumda?"}
    ).json()
    assert data["message"] == "DEMO-1 ne durumda?"


def test_teams_link_can_pick_a_template(api_client, conn):
    setup_group(api_client, conn)
    add_contacts(api_client)
    second = api_client.get("/api/templates").json()["templates"][1]
    data = api_client.post(
        "/api/issues/DEMO-1/teams-link", json={"template_id": second["id"]}
    ).json()
    assert data["message"].startswith("DEMO-1 için kısa bir güncelleme")


def test_people_mode_without_contacts_is_refused(api_client, conn):
    setup_group(api_client, conn)
    response = api_client.post("/api/issues/DEMO-1/teams-link", json={})
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "no_contacts"


def test_teams_link_for_an_unknown_issue_is_404(api_client, conn):
    setup_group(api_client, conn)
    response = api_client.post("/api/issues/DEMO-404/teams-link", json={})
    assert response.status_code == 404


def test_teams_link_without_any_template_says_so(api_client, conn):
    setup_group(api_client, conn)
    add_contacts(api_client)
    for template in api_client.get("/api/templates").json()["templates"]:
        api_client.delete(f"/api/templates/{template['id']}")
    response = api_client.post("/api/issues/DEMO-1/teams-link", json={})
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "template_missing"


def test_local_field_values_reach_the_message(api_client, conn):
    setup_group(api_client, conn)
    add_contacts(api_client)
    field = api_client.post(
        "/api/local-fields", json={"name": "Müşteri durumu", "type": "text"}
    ).json()["field"]
    api_client.put(f"/api/issues/DEMO-1/local/{field['id']}", json={"value": "müşteriye soruldu"})

    data = api_client.post(
        "/api/issues/DEMO-1/teams-link",
        json={"body": "Durum: {local:%d}" % field["id"]},
    ).json()
    assert data["message"] == "Durum: müşteriye soruldu"


def test_message_preview_does_not_record_anything(api_client, conn):
    setup_group(api_client, conn)
    add_contacts(api_client)
    preview = api_client.post("/api/issues/DEMO-1/message-preview", json={}).json()
    assert preview["message"].startswith("Merhaba, DEMO-1")
    assert api_client.get("/api/issues/DEMO-1/messages").json()["messages"] == []


def test_a_long_message_is_trimmed_and_the_full_text_goes_to_the_clipboard(api_client, conn):
    setup_group(api_client, conn)
    add_contacts(api_client)
    long_body = "ç" * 3000
    data = api_client.post("/api/issues/DEMO-1/teams-link", json={"body": long_body}).json()
    assert data["truncated"] is True
    assert data["clipboard"] == long_body
    assert len(data["url"]) <= teams.URL_LIMIT
    # Kayit notu tam metni tutar, kirpilmis hali degil.
    assert api_client.get("/api/issues/DEMO-1/messages").json()["messages"][0]["body"] == long_body


# --- sanal sutun --------------------------------------------------------


def test_contacts_column_is_offered_in_the_field_picker(api_client, conn):
    setup_group(api_client, conn)
    fields = api_client.get("/api/fields").json()["fields"]
    found = [item for item in fields if item["id"] == teams.CONTACTS_COLUMN]
    assert found and found[0]["name"] == "Teams kişileri"
    assert found[0]["virtual"] is True


def test_contacts_column_shows_names_in_the_grid(api_client, conn):
    group = setup_group(api_client, conn)
    add_contacts(
        api_client, contacts=[{"email": ORNEK, "name": "Örnek Kişi"}, IKINCI]
    )
    api_client.put(
        f"/api/groups/{group['id']}",
        json={"columns": ["issuekey", teams.CONTACTS_COLUMN]},
    )
    grid = api_client.get(f"/api/groups/{group['id']}/issues").json()
    assert [head["name"] for head in grid["columns"]] == ["Anahtar", "Teams kişileri"]
    cell = grid["rows"][0]["cells"][1]
    # Ad varsa ad, yoksa e-posta.
    assert cell["text"] == f"Örnek Kişi, {IKINCI}"


def test_contacts_column_can_be_sorted(api_client, conn):
    group = setup_group(api_client, conn, keys=("DEMO-1", "DEMO-2"))
    add_contacts(api_client, "DEMO-1", [{"email": ORNEK, "name": "Zeynep"}])
    add_contacts(api_client, "DEMO-2", [{"email": IKINCI, "name": "Ahmet"}])
    api_client.put(
        f"/api/groups/{group['id']}",
        json={"columns": ["issuekey", teams.CONTACTS_COLUMN]},
    )
    rows = api_client.get(
        f"/api/groups/{group['id']}/issues",
        params={"sort": teams.CONTACTS_COLUMN, "dir": "asc"},
    ).json()["rows"]
    assert [row["key"] for row in rows] == ["DEMO-2", "DEMO-1"]


def test_contacts_column_lands_in_excel(api_client, conn):
    group = setup_group(api_client, conn)
    add_contacts(api_client, contacts=[{"email": ORNEK, "name": "Örnek Kişi"}])
    response = api_client.get(
        f"/api/groups/{group['id']}/export.xlsx",
        params={"columns": f"issuekey,{teams.CONTACTS_COLUMN}"},
    )
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).worksheets[0]
    assert [cell.value for cell in sheet[1]] == ["Anahtar", "Teams kişileri"]
    assert sheet.cell(row=2, column=2).value == "Örnek Kişi"


def test_removing_a_contact_empties_the_column(api_client, conn):
    group = setup_group(api_client, conn)
    add_contacts(api_client)
    add_contacts(api_client, contacts=[])
    api_client.put(f"/api/groups/{group['id']}", json={"columns": [teams.CONTACTS_COLUMN]})
    grid = api_client.get(f"/api/groups/{group['id']}/issues").json()
    assert grid["rows"][0]["cells"][0]["text"] == ""


# --- arayuz kancalari ---------------------------------------------------


def test_drawer_carries_the_teams_section(api_client):
    script = api_client.get("/static/js/app.js").text
    for marker in (
        "renderDrawerTeams",
        "sendTeamsLink",
        "askStatus",
        '"Teams\'te aç"',
        # Acmayi sunucu yapar; sekme yalnizca yedek yolda acilir.
        "teams-open",
        "message-preview",
    ):
        assert marker in script, marker
    # Kirpilan mesajin tamami panoya gider; kanal yolu yok.
    assert "tamamı panoya kopyalandı" in script
    assert "Teams'te açıldı" in script
    # Sekme yalnizca sunucu acamadiysa acilir.
    assert "if (!data.opened) window.open(" in script
    assert "channel" not in script
    common = api_client.get("/static/js/common.js").text
    assert "copyText" in common and "navigator.clipboard" in common


def test_settings_page_has_the_teams_card(api_client):
    page = api_client.get("/settings").text
    for marker in (
        'id="teams-card"',
        'id="teams-templates"',
        'id="teams-contacts"',
        'id="teams-topic"',
        'id="teams-gal"',
        "Rehberi Outlook'tan yenile",
        "Mesaj Teams'te hazır açılır, göndermek için Gönder'e basmanız gerekir.",
    ):
        assert marker in page, marker
    assert "kanal" not in page.lower()
    script = api_client.get("/static/js/settings.js").text
    for marker in ("loadTeamsLists", "addTemplate", "dropContact", "importGal"):
        assert marker in script, marker


def test_teams_styles_are_in_the_sheet(api_client):
    css = api_client.get("/static/css/app.css").text
    for marker in (".drawer-teams", ".teams-chip", ".teams-sent", ".sent-line", ".teams-row"):
        assert marker in css, marker


def test_the_chat_icon_exists(api_client):
    common = api_client.get("/static/js/common.js").text
    assert "chat: [" in common
    script = api_client.get("/static/js/app.js").text
    assert 'icon("chat")' in script


# --- kurum rehberi (GAL) ------------------------------------------------
#
# Sahte COM nesneleri: gercek Outlook'a hic dokunulmaz. Rehberde odalar ve
# X500'de kalmis eski kayitlar da bulunur; cozulemeyen giris atlanmali.


class GalUser:
    def __init__(self, address):
        self.PrimarySmtpAddress = address


class GalAccessor:
    def __init__(self, values):
        self._values = values

    def GetProperty(self, tag):  # noqa: N802 - COM adi
        if tag not in self._values:
            raise RuntimeError("property not found")
        return self._values[tag]


class GalAddressEntry:
    """`AddressEntry`: Exchange kullanicisi, dagitim listesi ya da hicbiri."""

    def __init__(self, name="", exchange=None, group=None, smtp=None, address="", kind=0):
        self.Name = name
        self.AddressEntryUserType = kind
        self.Address = address
        self._exchange = exchange
        self._group = group
        self.PropertyAccessor = GalAccessor(
            {outlook.PR_SMTP_ADDRESS: smtp} if smtp else {}
        )

    def GetExchangeUser(self):  # noqa: N802 - COM adi
        return GalUser(self._exchange) if self._exchange else None

    def GetExchangeDistributionList(self):  # noqa: N802 - COM adi
        return GalUser(self._group) if self._group else None


class GalEntries:
    """`AddressEntries`: yalnizca GetFirst/GetNext ile yurunur (tek gecis)."""

    def __init__(self, entries):
        self._entries = list(entries)
        self._index = 0
        self.walks = 0

    def GetFirst(self):  # noqa: N802 - COM adi
        self.walks += 1
        self._index = 0
        return self._next()

    def GetNext(self):  # noqa: N802 - COM adi
        return self._next()

    def _next(self):
        if self._index >= len(self._entries):
            return None
        item = self._entries[self._index]
        self._index += 1
        return item

    @property
    def Count(self):  # noqa: N802 - COM adi
        raise RuntimeError("buyuk rehberde Count okunmamali")


class GalList:
    def __init__(self, entries):
        self.AddressEntries = GalEntries(entries)


class GalNamespace:
    def __init__(self, entries):
        self.gal = GalList(entries)

    def GetGlobalAddressList(self):  # noqa: N802 - COM adi
        return self.gal


class NoGalNamespace:
    """Exchange'siz kurulum: rehber cagirisi hic yok."""


@pytest.fixture
def outlook_source(monkeypatch):
    monkeypatch.setattr(outlook.sys, "platform", "win32")
    return outlook.OutlookSource()


def test_gal_resolves_smtp_through_three_steps(outlook_source):
    namespace = GalNamespace(
        [
            GalAddressEntry(name="Ayşe Yılmaz", exchange="ayse@example.com"),
            GalAddressEntry(name="Mehmet Demir", smtp="mehmet@example.com",
                            address="/O=KURUM/CN=MEHMET"),
            GalAddressEntry(name="Dış Kişi", address="disari@example.com"),
        ]
    )
    entries = outlook_source._address_book(namespace)
    assert [(item.name, item.email) for item in entries] == [
        ("Ayşe Yılmaz", "ayse@example.com"),
        ("Mehmet Demir", "mehmet@example.com"),
        ("Dış Kişi", "disari@example.com"),
    ]
    assert all(item.kind == "person" for item in entries)


def test_gal_marks_distribution_lists(outlook_source):
    namespace = GalNamespace(
        [
            GalAddressEntry(name="Tedarik Ekibi", group="tedarik@example.com",
                            kind=outlook.OL_EXCHANGE_DISTRIBUTION_LIST),
            GalAddressEntry(name="Ali", exchange="ali@example.com"),
        ]
    )
    entries = outlook_source._address_book(namespace)
    assert [(item.email, item.kind) for item in entries] == [
        ("tedarik@example.com", "list"),
        ("ali@example.com", "person"),
    ]


def test_gal_skips_entries_without_an_address(outlook_source):
    namespace = GalNamespace(
        [
            GalAddressEntry(name="Toplantı Odası 3", address="/O=KURUM/CN=ODA3"),
            GalAddressEntry(name="Boş"),
            GalAddressEntry(name="Ali", exchange="ali@example.com"),
        ]
    )
    assert [item.email for item in outlook_source._address_book(namespace)] == [
        "ali@example.com"
    ]


def test_gal_is_walked_once_with_getfirst_getnext(outlook_source):
    namespace = GalNamespace([GalAddressEntry(name=f"K{i}", exchange=f"k{i}@example.com")
                              for i in range(50)])
    entries = outlook_source._address_book(namespace)
    assert len(entries) == 50
    # Tek gecis: imlec bir kez basa alinir, `Count` hic okunmaz (okunsa patlardi).
    assert namespace.gal.AddressEntries.walks == 1


def test_gal_keeps_the_first_of_a_repeated_address(outlook_source):
    namespace = GalNamespace(
        [
            GalAddressEntry(name="Ali Can", exchange="ali@example.com"),
            GalAddressEntry(name="Ali C.", exchange="ALI@example.com"),
        ]
    )
    entries = outlook_source._address_book(namespace)
    assert [(item.name, item.email) for item in entries] == [("Ali Can", "ali@example.com")]


def test_gal_without_exchange_says_so(outlook_source):
    with pytest.raises(MailError) as caught:
        outlook_source._address_book(NoGalNamespace())
    assert caught.value.code == "gal_unavailable"


def test_import_gal_writes_the_address_book(api_client, fake_mail):
    result = api_client.post("/api/contacts/import-gal").json()["result"]
    assert result["imported"] == 3
    assert result["updated"] == 0
    assert result["lists"] == 1
    assert result["total"] == 3
    assert result["ms"] >= 0
    assert fake_mail.gal_reads == 1

    book = api_client.get("/api/contacts").json()["contacts"]
    assert [item["email"] for item in book] == [
        "ayse.yilmaz@example.com",
        "mehmet.demir@example.com",
        "tedarik@example.com",
    ]
    assert [item["source"] for item in book] == ["gal", "gal", "gal"]
    assert [item["kind"] for item in book] == ["person", "person", "list"]


def test_import_gal_updates_names_on_the_second_run(api_client, fake_mail):
    api_client.post("/api/contacts/import-gal")
    fake_mail.gal = [
        GalEntry(email="ayse.yilmaz@example.com", name="Ayşe Yılmaz (Satın Alma)"),
    ]
    result = api_client.post("/api/contacts/import-gal").json()["result"]
    assert result == {
        "imported": 0,
        "updated": 1,
        "lists": 0,
        "total": 1,
        "ms": result["ms"],
        "synced_at": result["synced_at"],
    }
    found = api_client.get("/api/contacts", params={"q": "ayse"}).json()["contacts"]
    assert found[0]["name"] == "Ayşe Yılmaz (Satın Alma)"
    # Ikinci tarama eskileri SILMEZ: baska kayitlarda kullaniliyor olabilir.
    assert len(api_client.get("/api/contacts").json()["contacts"]) == 3


def test_import_gal_keeps_hand_written_contacts(api_client, conn, fake_mail):
    setup_group(api_client, conn)
    add_contacts(api_client, contacts=[{"email": "elle@example.com", "name": "Elle Yazılan"}])
    fake_mail.gal = [GalEntry(email="elle@example.com", name="Rehberdeki Ad")]

    result = api_client.post("/api/contacts/import-gal").json()["result"]
    assert result["imported"] == 0 and result["updated"] == 1

    book = {item["email"]: item for item in api_client.get("/api/contacts").json()["contacts"]}
    # Kaynak "manual" kalir (silinmesin diye), ad rehberden tazelenir.
    assert book["elle@example.com"]["source"] == "manual"
    assert book["elle@example.com"]["name"] == "Rehberdeki Ad"
    # Kayittaki bag da yerinde.
    assert api_client.get("/api/issues/DEMO-1/contacts").json()["contacts"][0]["email"] == (
        "elle@example.com"
    )


def test_import_gal_skips_broken_addresses(api_client, fake_mail):
    fake_mail.gal = [
        GalEntry(email="saglam@example.com", name="Sağlam"),
        GalEntry(email="bozuk", name="Bozuk"),
        GalEntry(email="", name="Boş"),
    ]
    result = api_client.post("/api/contacts/import-gal").json()["result"]
    assert result["total"] == 1 and result["imported"] == 1
    assert [item["email"] for item in api_client.get("/api/contacts").json()["contacts"]] == [
        "saglam@example.com"
    ]


def test_import_gal_stamps_the_last_sync(api_client):
    assert api_client.get("/api/settings").json()["settings"]["teams.gal_synced_at"] == ""
    result = api_client.post("/api/contacts/import-gal").json()["result"]
    stamped = api_client.get("/api/settings").json()["settings"]["teams.gal_synced_at"]
    assert stamped and stamped == result["synced_at"]


def test_gal_import_is_unavailable_without_outlook(api_client, context):
    def missing(body_limit=None):
        raise mail_unavailable()

    context.mail_factory = missing
    response = api_client.post("/api/contacts/import-gal")
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "feature_unavailable"


def test_search_mixes_the_book_and_the_directory(api_client, conn, fake_mail):
    setup_group(api_client, conn)
    add_contacts(api_client, contacts=[{"email": "ayhan@example.com", "name": "Ayhan Elle"}])
    fake_mail.gal = [
        GalEntry(email="ayse.yilmaz@example.com", name="Ayşe Yılmaz"),
        GalEntry(email="kaya.ay@example.com", name="Kaya Ay"),
    ]
    api_client.post("/api/contacts/import-gal")

    found = api_client.get("/api/contacts", params={"q": "ay"}).json()["contacts"]
    # Onek eslesmeleri once (ad ya da e-posta bastan tutuyorsa), sonra icerenler.
    assert [item["name"] for item in found] == ["Ayhan Elle", "Ayşe Yılmaz", "Kaya Ay"]
    assert {item["source"] for item in found} == {"manual", "gal"}


def test_directory_lists_can_be_attached_to_an_issue(api_client, conn, fake_mail):
    setup_group(api_client, conn)
    api_client.post("/api/contacts/import-gal")
    saved = add_contacts(api_client, contacts=["tedarik@example.com"])
    # Dagitim listesi oldugu arayuze de gecer (rozet icin).
    assert saved == [{"email": "tedarik@example.com", "name": "Tedarik Ekibi", "kind": "list"}]


# --- sunucudan acma (msteams:) ------------------------------------------
#
# Tarayicidan acmak geride bos bir sekme birakiyordu; adres artik isletim
# sistemine veriliyor. Testlerde hicbir sey gercekten acilmaz.


@pytest.fixture
def opener(monkeypatch):
    """`desktop.open_url` yerine sayac: ne cagrildi, ne dondu."""
    calls: list[str] = []

    def fake(url, calls=calls):
        calls.append(url)
        return fake.result

    fake.result = True
    monkeypatch.setattr(desktop, "open_url", fake)
    fake.calls = calls
    return fake


def test_app_link_uses_the_msteams_protocol():
    link = teams.build_app_link([ORNEK], "Merhaba")
    assert link["url"].startswith("msteams:/l/chat/0/0?users=")
    assert link["url"] == f"msteams:/l/chat/0/0?users={ORNEK}&message=Merhaba"
    assert link["truncated"] is False


def test_app_link_carries_the_same_parameters_as_the_web_link():
    app_link = teams.build_app_link([ORNEK, IKINCI], "Durum?", topic="DEMO-1")
    web_link = teams.build_chat_link([ORNEK, IKINCI], "Durum?", topic="DEMO-1")
    assert app_link["url"].split("?", 1)[1] == web_link["url"].split("?", 1)[1]
    assert query_of(app_link["url"])["topicName"] == ["DEMO-1"]


def test_both_links_trim_to_the_very_same_text():
    long_text = "ç" * 4000
    app_link = teams.build_app_link([ORNEK], long_text)
    web_link = teams.build_chat_link([ORNEK], long_text)
    # Kirpma her zaman daha uzun olan https tabanina gore hesaplanir.
    assert app_link["message"] == web_link["message"]
    assert app_link["truncated"] and web_link["truncated"]
    assert len(app_link["url"]) <= teams.URL_LIMIT


def test_teams_open_hands_the_protocol_link_to_the_system(api_client, conn, opener):
    setup_group(api_client, conn)
    add_contacts(api_client)

    data = api_client.post("/api/issues/DEMO-1/teams-open", json={}).json()
    assert data["opened"] is True
    assert data["app_url"].startswith("msteams:/l/chat/0/0?users=")
    # Yedek adres yine yaninda: acilamasaydi arayuz sekme acacakti.
    assert data["url"].startswith("https://teams.microsoft.com/")
    assert opener.calls == [data["app_url"]]

    history = api_client.get("/api/issues/DEMO-1/messages").json()["messages"]
    assert len(history) == 1
    assert history[0]["body"] == data["message"]


def test_teams_open_reports_failure_so_the_browser_can_take_over(api_client, conn, opener):
    setup_group(api_client, conn)
    add_contacts(api_client)
    opener.result = False

    data = api_client.post("/api/issues/DEMO-1/teams-open", json={}).json()
    assert data["opened"] is False
    assert query_of(data["url"])["users"] == [ORNEK]
    # Acilamasa da "acildi" kaydi bir kez yazilir: kullanici yine gonderecek.
    assert len(api_client.get("/api/issues/DEMO-1/messages").json()["messages"]) == 1


def test_teams_open_accepts_a_template_or_a_body(api_client, conn, opener):
    setup_group(api_client, conn)
    add_contacts(api_client)
    second = api_client.get("/api/templates").json()["templates"][1]

    picked = api_client.post(
        "/api/issues/DEMO-1/teams-open", json={"template_id": second["id"]}
    ).json()
    assert picked["message"].startswith("DEMO-1 için kısa bir güncelleme")

    written = api_client.post(
        "/api/issues/DEMO-1/teams-open", json={"body": "{key} ne durumda?"}
    ).json()
    assert written["message"] == "DEMO-1 ne durumda?"


def test_teams_open_without_contacts_is_refused_and_opens_nothing(api_client, conn, opener):
    setup_group(api_client, conn)
    response = api_client.post("/api/issues/DEMO-1/teams-open", json={})
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "no_contacts"
    assert opener.calls == []


def test_teams_open_for_an_unknown_issue_is_404(api_client, conn, opener):
    setup_group(api_client, conn)
    assert api_client.post("/api/issues/DEMO-404/teams-open", json={}).status_code == 404
    assert opener.calls == []


def test_teams_link_still_works_and_opens_nothing(api_client, conn, opener):
    """Onizleme/yedek yolu duruyor: yalnizca adresi doner, hicbir sey acmaz."""
    setup_group(api_client, conn)
    add_contacts(api_client)
    data = api_client.post("/api/issues/DEMO-1/teams-link", json={}).json()
    assert data["url"].startswith("https://teams.microsoft.com/")
    assert "app_url" not in data and "opened" not in data
    assert opener.calls == []


def test_long_message_is_copied_and_trimmed_on_the_open_path(api_client, conn, opener):
    setup_group(api_client, conn)
    add_contacts(api_client)
    long_body = "ç" * 3000
    data = api_client.post("/api/issues/DEMO-1/teams-open", json={"body": long_body}).json()
    assert data["truncated"] is True
    assert data["clipboard"] == long_body
    assert len(data["app_url"]) <= teams.URL_LIMIT


# --- adresi isletim sistemine verme -------------------------------------


def test_open_url_uses_startfile_on_windows(monkeypatch):
    opened: list[str] = []
    monkeypatch.setattr(desktop.sys, "platform", "win32")
    monkeypatch.setattr(desktop.os, "startfile", opened.append, raising=False)
    assert desktop.open_url("msteams:/l/chat/0/0?users=a@example.com") is True
    assert opened == ["msteams:/l/chat/0/0?users=a@example.com"]


def test_open_url_survives_a_windows_failure(monkeypatch):
    def angry(url):
        raise OSError("kayitli isleyici yok")

    monkeypatch.setattr(desktop.sys, "platform", "win32")
    monkeypatch.setattr(desktop.os, "startfile", angry, raising=False)
    # Hata disari sizmaz: arayuz `opened: false` gorup sekmeye duser.
    assert desktop.open_url("msteams:/l/chat") is False


def test_open_url_uses_the_desktop_helper_elsewhere(monkeypatch):
    calls: list[list[str]] = []
    monkeypatch.setattr(desktop.subprocess, "Popen", lambda command, **kw: calls.append(command))

    monkeypatch.setattr(desktop.sys, "platform", "linux")
    assert desktop.open_url("msteams:/l/chat") is True
    monkeypatch.setattr(desktop.sys, "platform", "darwin")
    assert desktop.open_url("msteams:/l/chat") is True
    assert calls == [["xdg-open", "msteams:/l/chat"], ["open", "msteams:/l/chat"]]


def test_open_url_without_a_helper_is_false(monkeypatch):
    def missing(command, **kw):
        raise FileNotFoundError(command[0])

    monkeypatch.setattr(desktop.sys, "platform", "linux")
    monkeypatch.setattr(desktop.subprocess, "Popen", missing)
    assert desktop.open_url("msteams:/l/chat") is False


def test_open_url_ignores_an_empty_address(monkeypatch):
    monkeypatch.setattr(desktop.sys, "platform", "linux")
    monkeypatch.setattr(
        desktop.subprocess, "Popen", lambda *a, **k: pytest.fail("boş adres açılmamalı")
    )
    assert desktop.open_url("   ") is False
